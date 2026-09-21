from __future__ import annotations

import argparse
import csv
import hashlib
import math
import os
import re
import sys
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path

from app.infrastructure.db.enums import PlaceCategory, PlaceSource

REQUIRED_COLUMNS = {
    "name",
    "type",
    "latitude",
    "longitude",
    "address",
    "phone",
    "phone 2",
    "instagram",
    "telegram",
    "opening_hours",
    "days_off",
    "website",
    "description",
}

TYPE_ALIASES = {
    "pet store": PlaceCategory.PET_SHOP,
    "pet shop": PlaceCategory.PET_SHOP,
    "pet_shop": PlaceCategory.PET_SHOP,
    "veterinary clinic": PlaceCategory.VETERINARY,
    "veterinary": PlaceCategory.VETERINARY,
    "vet": PlaceCategory.VETERINARY,
    "animal shelter": PlaceCategory.SHELTER,
    "shelter": PlaceCategory.SHELTER,
}


@dataclass(frozen=True, slots=True)
class ManualPlace:
    row_number: int
    source_id: str
    name: str
    category: PlaceCategory
    categories: tuple[PlaceCategory, ...]
    latitude: float
    longitude: float
    address: str | None
    phone: str | None
    phone_2: str | None
    instagram: str | None
    telegram: str | None
    opening_hours: str | None
    days_off: str | None
    website: str | None
    description: str | None


@dataclass(slots=True)
class ImportIssue:
    row_number: int
    message: str


@dataclass(slots=True)
class ImportReport:
    rows_processed: int = 0
    places_created: int = 0
    duplicates_skipped: int = 0
    invalid_rows: list[ImportIssue] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def read_rows(path: Path) -> list[tuple[int, dict[str, object]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv_rows(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_excel_rows(path)
    raise ValueError("Unsupported file type. Use .xlsx, .xlsm, or .csv.")


def parse_manual_place_row(row_number: int, row: dict[str, object]) -> ManualPlace:
    name = _required_text(row.get("name"), "name")
    categories = _parse_categories(row.get("type"))
    category = _primary_category(categories)
    latitude = _parse_coordinate(row.get("latitude"), "latitude", -90, 90)
    longitude = _parse_coordinate(row.get("longitude"), "longitude", -180, 180)
    source_id = _source_id(
        name=name,
        categories=categories,
        latitude=latitude,
        longitude=longitude,
    )
    return ManualPlace(
        row_number=row_number,
        source_id=source_id,
        name=name,
        category=category,
        categories=categories,
        latitude=latitude,
        longitude=longitude,
        address=_clean_text(row.get("address")),
        phone=_clean_text(row.get("phone")),
        phone_2=_clean_text(row.get("phone 2")),
        instagram=_clean_text(row.get("instagram")),
        telegram=_clean_text(row.get("telegram")),
        opening_hours=_clean_text(row.get("opening_hours")),
        days_off=_clean_text(row.get("days_off")),
        website=_clean_text(row.get("website")),
        description=_clean_text(row.get("description")),
    )


def parse_manual_places(
    rows: Iterable[tuple[int, dict[str, object]]],
) -> tuple[list[ManualPlace], ImportReport]:
    places: list[ManualPlace] = []
    report = ImportReport()
    seen_source_ids: set[str] = set()
    for row_number, row in rows:
        report.rows_processed += 1
        try:
            place = parse_manual_place_row(row_number, row)
        except ValueError as exc:
            report.invalid_rows.append(ImportIssue(row_number=row_number, message=str(exc)))
            continue
        if place.source_id in seen_source_ids:
            report.duplicates_skipped += 1
            continue
        seen_source_ids.add(place.source_id)
        places.append(place)
    return places, report


def import_places(
    session,
    places: Iterable[ManualPlace],
    report: ImportReport,
) -> ImportReport:
    from geoalchemy2.elements import WKTElement
    from sqlalchemy import select

    from app.infrastructure.db.models import schema

    for place in places:
        existing = session.scalars(
            select(schema.Place).where(
                schema.Place.source == PlaceSource.MANUAL,
                schema.Place.source_id == place.source_id,
            )
        ).one_or_none()
        if existing is not None:
            report.duplicates_skipped += 1
            continue

        place_record = schema.Place(
            name=place.name,
            category=place.category,
            location=WKTElement(f"POINT({place.longitude} {place.latitude})", srid=4326),
            address=place.address,
            phone=place.phone,
            phone_2=place.phone_2,
            instagram=place.instagram,
            telegram=place.telegram,
            opening_hours=place.opening_hours,
            days_off=place.days_off,
            website=place.website,
            description=place.description,
            source=PlaceSource.MANUAL,
            source_id=place.source_id,
            is_active=True,
        )
        place_record.category_links = [
            schema.PlaceCategoryLink(category=category) for category in place.categories
        ]
        session.add(place_record)
        report.places_created += 1
    session.flush()
    return report


def _read_csv_rows(path: Path) -> list[tuple[int, dict[str, object]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        _validate_columns(reader.fieldnames or [])
        return [(reader.line_num, dict(row)) for row in reader]


def _read_excel_rows(path: Path) -> list[tuple[int, dict[str, object]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise ValueError("Import file is empty.") from None

    columns = [_clean_header(value) for value in header]
    _validate_columns(columns)
    result: list[tuple[int, dict[str, object]]] = []
    for offset, values in enumerate(rows, start=2):
        if values is None or all(_clean_text(value) is None for value in values):
            continue
        row = {
            column: values[index] if index < len(values) else None
            for index, column in enumerate(columns)
        }
        result.append((offset, row))
    return result


def _validate_columns(columns: Iterable[str]) -> None:
    normalized = {_clean_header(column) for column in columns}
    missing = REQUIRED_COLUMNS - normalized
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(sorted(missing)))


def _parse_categories(value: object) -> tuple[PlaceCategory, ...]:
    cleaned = _clean_text(value)
    if cleaned is None:
        raise ValueError("type is required")
    categories: list[PlaceCategory] = []
    unsupported: list[str] = []
    for part in cleaned.split(","):
        normalized = part.lower().replace("-", " ")
        normalized = re.sub(r"\s+", " ", normalized).strip()
        category = TYPE_ALIASES.get(normalized)
        if category is None:
            unsupported.append(part.strip())
            continue
        if category not in categories:
            categories.append(category)
    if unsupported:
        raise ValueError(f"unsupported type: {cleaned}")
    return tuple(sorted(categories, key=_category_priority))


def _parse_coordinate(value: object, field_name: str, minimum: float, maximum: float) -> float:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{field_name} is required")
    if isinstance(value, int | float):
        coordinate = float(value)
    else:
        cleaned = _clean_text(value)
        if cleaned is None:
            raise ValueError(f"{field_name} is required")
        try:
            coordinate = float(cleaned.replace(",", "."))
        except ValueError as exc:
            raise ValueError(f"{field_name} must be a number") from exc
    if not math.isfinite(coordinate) or coordinate < minimum or coordinate > maximum:
        raise ValueError(f"{field_name} is out of range")
    return coordinate


def _required_text(value: object, field_name: str) -> str:
    cleaned = _clean_text(value)
    if cleaned is None:
        raise ValueError(f"{field_name} is required")
    return cleaned


def _clean_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    text = str(value).replace("\xa0", " ").strip()
    text = re.sub(r"\s+", " ", text)
    if not text or text.lower() == "nan":
        return None
    return text


def _clean_header(value: object) -> str:
    cleaned = _clean_text(value)
    return "" if cleaned is None else cleaned.lower()


def _source_id(
    *,
    name: str,
    categories: tuple[PlaceCategory, ...],
    latitude: float,
    longitude: float,
) -> str:
    fingerprint = "|".join(
        [
            name.casefold(),
            ",".join(category.value for category in categories),
            f"{latitude:.6f}",
            f"{longitude:.6f}",
        ]
    )
    digest = hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()[:24]
    return f"manual_file:{digest}"


def _primary_category(categories: tuple[PlaceCategory, ...]) -> PlaceCategory:
    return sorted(categories, key=_category_priority)[0]


def _category_priority(category: PlaceCategory) -> int:
    priority = {
        PlaceCategory.VETERINARY: 0,
        PlaceCategory.SHELTER: 1,
        PlaceCategory.PET_SHOP: 2,
    }
    return priority[category]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Import manually prepared places from Excel or CSV."
    )
    parser.add_argument("file", type=Path, help="Path to .xlsx, .xlsm, or .csv import file.")
    parser.add_argument("--database-url", default=os.getenv("DATABASE_URL"))
    parser.add_argument("--dry-run", action="store_true")
    return parser.parse_args()


def print_report(report: ImportReport) -> None:
    print(f"rows processed: {report.rows_processed}")
    print(f"places created: {report.places_created}")
    print(f"duplicates skipped: {report.duplicates_skipped}")
    print(f"invalid rows: {len(report.invalid_rows)}")
    print(f"errors: {len(report.errors)}")
    for issue in report.invalid_rows[:20]:
        print(f"invalid row {issue.row_number}: {issue.message}")
    if len(report.invalid_rows) > 20:
        print(f"... {len(report.invalid_rows) - 20} more invalid rows")
    for error in report.errors:
        print(f"error: {error}", file=sys.stderr)


def main() -> int:
    args = parse_args()
    if not args.file.exists():
        print(f"Import file not found: {args.file}", file=sys.stderr)
        return 2
    if not args.database_url and not args.dry_run:
        print("DATABASE_URL is required unless --dry-run is used.", file=sys.stderr)
        return 2

    report = ImportReport()
    try:
        rows = read_rows(args.file)
        places, report = parse_manual_places(rows)
    except Exception as exc:
        report.errors.append(str(exc))
        print_report(report)
        return 1

    if args.dry_run:
        print_report(report)
        return 0 if not report.invalid_rows else 1

    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session, sessionmaker

    engine = create_engine(args.database_url, pool_pre_ping=True)
    session_factory = sessionmaker(bind=engine, class_=Session, expire_on_commit=False)
    try:
        with session_factory.begin() as session:
            import_places(session, places, report)
    except Exception as exc:
        report.errors.append(str(exc))
        print_report(report)
        return 1

    print_report(report)
    return 0 if not report.invalid_rows and not report.errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
