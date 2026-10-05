"""Present the saved, read-only final Places V2 deployment diff for human review."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "places_v2" / "output"
WORKBOOK = OUT / "production_deployment_review_final.xlsx"


def read(name: str) -> list[dict]:
    return [json.loads(line) for line in (OUT / name).read_text(encoding="utf-8").splitlines()]


def sheet(book: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    page = book.create_sheet(title)
    page.append(headers)
    for row in rows:
        page.append(row)
    page.freeze_panes = "D2"
    page.auto_filter.ref = page.dimensions
    for cell in page[1]:
        cell.fill = PatternFill("solid", fgColor="244268")
        cell.font = Font(color="FFFFFF", bold=True)
    for column in ("A", "B", "C"):
        page.column_dimensions[column].width = 33


def main() -> None:
    actions = read("production_deployment_review_final.jsonl")
    pairs = read("production_pair_decisions_final.jsonl")
    fields = read("production_field_decisions_final.jsonl")
    report = json.loads(
        (OUT / "production_deployment_review_report_final.json").read_text(encoding="utf-8")
    )
    book = Workbook()
    book.remove(book.active)
    sheet(
        book,
        "Summary",
        ["metric", "value"],
        [
            [
                key,
                json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value,
            ]
            for key, value in report.items()
        ],
    )
    sheet(
        book,
        "Actions",
        [
            "action",
            "entity",
            "production_uuid",
            "production_name",
            "v2_id",
            "v2_name",
            "production_category",
            "v2_category",
            "production_address",
            "v2_address",
            "source_id",
            "new_uuid",
            "category_links_before",
            "category_links_after",
            "reason",
            "approve_notes",
        ],
        [
            [
                row["action"],
                row["entity"],
                row.get("production_uuid"),
                row.get("production_name"),
                row.get("v2_id"),
                row.get("v2_name"),
                row.get("production_category"),
                row.get("v2_category"),
                row.get("production_address"),
                row.get("v2_address"),
                row.get("v2_source_id"),
                row.get("new_uuid"),
                ", ".join(row.get("category_links_before", [])),
                ", ".join(row.get("category_links_after", [])),
                row.get("reason"),
                None,
            ]
            for row in actions
        ],
    )
    sheet(
        book,
        "69 pair decisions",
        [
            "pair_number",
            "decision",
            "actionable_mapping",
            "production_uuid",
            "production_name",
            "production_category",
            "production_address",
            "v2_id",
            "v2_name",
            "v2_category",
            "v2_address",
            "distance_m",
            "name_similarity",
            "address_similarity",
            "phone_agreement",
            "website_agreement",
            "decision_reason",
            "source_urls",
            "review_notes",
        ],
        [
            [
                row["pair_number"],
                row["decision"],
                row.get("actionable_mapping"),
                row["production_uuid"],
                row["production_name"],
                row["production_category"],
                row["production_address"],
                row["v2_id"],
                row["v2_name"],
                row["v2_category"],
                row["v2_address"],
                row["distance_m"],
                row["matching_evidence"]["name_similarity"],
                row["matching_evidence"]["address_similarity"],
                row["matching_evidence"]["phone_agreement"],
                row["matching_evidence"]["website_agreement"],
                row["decision_reason"],
                "\n".join(row["v2_evidence_urls"]),
                None,
            ]
            for row in pairs
        ],
    )
    sheet(
        book,
        "Field operations",
        [
            "production_uuid",
            "v2_id",
            "field",
            "operation",
            "production_value",
            "v2_value",
            "result_value",
            "v2_confidence",
            "evidence_urls",
            "reason",
            "review_notes",
        ],
        [
            [
                row["production_uuid"],
                row["v2_id"],
                row["field"],
                row["operation"],
                str(row["production_value"]) if row["production_value"] is not None else None,
                str(row["v2_value"]) if row["v2_value"] is not None else None,
                str(row["result_value"]) if row["result_value"] is not None else None,
                row["v2_confidence"],
                "\n".join(row["evidence_urls"]),
                row["reason"],
                None,
            ]
            for row in fields
        ],
    )
    sheet(
        book,
        "Remaining decisions",
        ["type", "production_uuid", "v2_id", "reason", "decision_notes"],
        [
            ["PAIR", row["production_uuid"], row["v2_id"], row["decision_reason"], None]
            for row in pairs
            if row["decision"] == "UNRESOLVED"
        ]
        + [
            [
                "DUPLICATE_LEGACY_UUID",
                row["production_uuid"],
                row["v2_id"],
                row["decision_reason"],
                None,
            ]
            for row in pairs
            if row.get("uuid_decision") == "SECOND_LEGACY_UUID_REVIEW"
        ]
        + [
            ["FIELD", row["production_uuid"], row["v2_id"], row["reason"], None]
            for row in fields
            if row["operation"] == "REVIEW"
        ],
    )
    book.save(WORKBOOK)
    print(
        json.dumps(
            {
                "workbook": str(WORKBOOK),
                "actions": len(actions),
                "pairs": len(pairs),
                "fields": len(fields),
            }
        )
    )


if __name__ == "__main__":
    main()
