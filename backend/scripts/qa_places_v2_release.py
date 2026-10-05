"""Offline, conservative QA for the reviewed V2 release candidate.

This reads only generated V2 research files. Warnings never merge or remove a
place. Run after build_places_v2.py; it adds a filterable Release candidate
sheet to the existing workbook and writes separate machine-readable QA files.
"""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from uuid import UUID, uuid5

from backend.scripts.build_places_v2 import (
    OUT,
    TASHKENT_BOUNDS,
    distance_meters,
    district_hint,
    normal,
    read_jsonl,
    write_jsonl,
)
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

NAMESPACE = UUID("b5bf982c-6927-42bf-9cd5-0a6ffdaf5a4e")
CATEGORY_MAP = {
    "pet_store": "pet_shop",
    "veterinary_clinic": "veterinary",
    "veterinary_pharmacy": "veterinary_pharmacy",
    "animal_shelter": "shelter",
}


def qa() -> dict:
    places = read_jsonl(OUT / "places_verified.jsonl")
    sources = read_jsonl(OUT / "places_sources.jsonl")
    existing_audit = read_jsonl(OUT / "places_audit.jsonl")
    warnings: list[dict] = []
    by_id: dict[str, list[dict]] = defaultdict(list)
    evidence: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for row in sources:
        evidence[row["record_id"], row["field"]].append(row)

    def warn(ids: list[str], code: str, level: str, detail: str) -> None:
        row = {"place_ids": ids, "code": code, "level": level, "detail": detail}
        warnings.append(row)
        for place_id in ids:
            by_id[place_id].append(row)

    ids = [p["id"] for p in places]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate V2 stable IDs")
    organization_urls: dict[str, list[str]] = defaultdict(list)
    for p in places:
        for url in p["source_urls"]:
            if "/firm/" in url or "/org/" in url:
                organization_urls[url].append(p["id"])
    for url, place_ids in organization_urls.items():
        if len(place_ids) > 1:
            warn(place_ids, "organization_url_reused", "review", url)
    for p in places:
        pid = p["id"]
        point = (p["latitude"], p["longitude"])
        firm_ids = {
            match.group(1)
            for url in p["source_urls"]
            if (match := re.search(r"2gis\.uz/tashkent/firm/(\d+)", url))
        }
        if len(firm_ids) > 1:
            warn(
                [pid],
                "multiple_firm_ids_one_place",
                "review",
                "Several 2GIS firm IDs attached; inspect whether chain branches were merged",
            )
        if p["category"] not in CATEGORY_MAP:
            warn([pid], "unsupported_category", "blocker", p["category"])
        for row in evidence[pid, "category"]:
            if row["status"] == "verified" and row["value"] != p["category"]:
                warn(
                    [pid],
                    "category_source_disagreement",
                    "review",
                    f"Record {p['category']}; source {row['value']}",
                )
        if (
            not p["name"]
            or len(p["name"].strip()) < 3
            or "\ufffd" in p["name"]
            or "РЎ" in p["name"]
        ):
            warn([pid], "malformed_name", "review", repr(p["name"]))
        if p["address"] and ("\ufffd" in p["address"] or "РЎ" in p["address"]):
            warn([pid], "malformed_address", "review", repr(p["address"]))
        if p["address"] and not any(c.isdigit() for c in p["address"]):
            warn(
                [pid],
                "address_without_house",
                "visual",
                "Inspect the source pin; street or mahallah alone is broad",
            )
        letters = [c for c in p["name"] if c.isalpha()]
        if any("a" <= c.casefold() <= "z" for c in letters) and any(
            "а" <= c.casefold() <= "я" for c in letters
        ):
            warn([pid], "mixed_script_name", "review", repr(p["name"]))
        if not all(isinstance(x, (int, float)) for x in point):
            warn([pid], "missing_coordinates", "blocker", "No numeric point")
            continue
        if not (
            TASHKENT_BOUNDS[0] <= point[0] <= TASHKENT_BOUNDS[1]
            and TASHKENT_BOUNDS[2] <= point[1] <= TASHKENT_BOUNDS[3]
        ):
            warn([pid], "outside_tashkent_bounds", "blocker", str(point))
        if not p["district"]:
            warn([pid], "missing_district", "review", "No district assigned")
        if p["address"] and re.search(r"\b(district|tumani|район)\b", p["address"], re.IGNORECASE):
            address_district = district_hint(p["address"])
            if address_district and p["district"] and address_district != p["district"]:
                warn(
                    [pid],
                    "address_district_mismatch",
                    "review",
                    f"Record {p['district']}; address {address_district}",
                )
        if not p["address"]:
            if p["field_confidence"].get("location_evidence") not in {"high", "medium"}:
                warn(
                    [pid],
                    "location_identity_insufficient",
                    "blocker",
                    "No address or alternate location evidence",
                )
            else:
                warn(
                    [pid],
                    "address_null_location_evidence",
                    "visual",
                    "Inspect the pin and source location description",
                )
        if not p["source_urls"] or p["entity_confidence"] != "high":
            warn(
                [pid],
                "identity_evidence_insufficient",
                "blocker",
                "Missing source or high entity confidence",
            )
        coord_rows = [
            r
            for r in evidence[pid, "coordinates"]
            if r["status"] == "verified" and isinstance(r["value"], list)
        ]
        if not coord_rows:
            warn([pid], "coordinate_source_missing", "blocker", "No field-level coordinate source")
        else:
            distances = [distance_meters(point, tuple(r["value"])) for r in coord_rows]
            if max(distances) > 150:
                warn(
                    [pid],
                    "coordinate_source_disagreement",
                    "review",
                    f"Largest source-point difference {max(distances):.0f} m",
                )
            elif max(distances) > 30:
                warn(
                    [pid],
                    "coordinate_source_offset",
                    "visual",
                    f"Largest source-point difference {max(distances):.0f} m",
                )
        for row in evidence[pid, "district"]:
            if (
                row["status"] == "verified"
                and row["value"]
                and normal(str(row["value"])) != normal(p["district"])
            ):
                warn(
                    [pid],
                    "district_source_disagreement",
                    "review",
                    f"Record {p['district']}; source {row['value']}",
                )
        if p["field_conflicts"]:
            warn([pid], "field_conflict", "review", ", ".join(p["field_conflicts"]))

    clusters: dict[str, set[str]] = defaultdict(set)
    for index, a in enumerate(places):
        nearest = sorted(
            (
                (
                    distance_meters(
                        (a["latitude"], a["longitude"]), (b["latitude"], b["longitude"])
                    ),
                    b,
                )
                for b in places
                if b["id"] != a["id"]
            ),
            key=lambda pair: pair[0],
        )[:5]
        if a["district"] and nearest and all(b["district"] != a["district"] for _, b in nearest):
            warn(
                [a["id"]],
                "district_spatial_outlier",
                "visual",
                f"No same-district place among five nearest; closest {nearest[0][0]:.0f} m",
            )
        for b in places[index + 1 :]:
            distance = distance_meters(
                (a["latitude"], a["longitude"]), (b["latitude"], b["longitude"])
            )
            name_same = normal(a["name"]) == normal(b["name"])
            address_same = bool(
                a["address"] and b["address"] and normal(a["address"]) == normal(b["address"])
            )
            if distance < 0.5:
                warn(
                    [a["id"], b["id"]],
                    "identical_coordinates",
                    "visual",
                    f"Same pin; names {a['name']} / {b['name']}",
                )
            if name_same and distance <= 150:
                warn(
                    [a["id"], b["id"]],
                    "same_name_nearby",
                    "review",
                    f"{distance:.1f} m; compare branches",
                )
                if a["category"] != b["category"]:
                    warn(
                        [a["id"], b["id"]],
                        "nearby_category_difference",
                        "review",
                        f"{a['category']} / {b['category']}; check one mixed-use branch",
                    )
            elif distance <= 20 and (not name_same):
                warn(
                    [a["id"], b["id"]],
                    "different_names_close",
                    "visual",
                    f"{distance:.1f} m; inspect shared building",
                )
            if address_same and not name_same:
                warn([a["id"], b["id"]], "same_address_different_names", "visual", a["address"])
            if distance <= 30:
                clusters[a["id"]].add(b["id"])
                clusters[b["id"]].add(a["id"])
    for pid, neighbors in clusters.items():
        if len(neighbors) >= 3:
            warn(
                [pid], "dense_coordinate_cluster", "visual", f"{len(neighbors)} places within 30 m"
            )

    rc = []
    for p in places:
        row = dict(p)
        row["source_id"] = "places_v2:" + p["id"]
        row["import_uuid"] = str(uuid5(NAMESPACE, row["source_id"]))
        row["include"] = "PENDING_HUMAN_REVIEW"
        row["qa_warnings"] = sorted({w["code"] for w in by_id[p["id"]]})
        row["qa_warning_count"] = len(by_id[p["id"]])
        rc.append(row)
    if len({r["source_id"] for r in rc}) != len(rc) or len({r["import_uuid"] for r in rc}) != len(
        rc
    ):
        raise ValueError("V2 import identity collision")
    coordinate_codes = {
        "identical_coordinates",
        "same_name_nearby",
        "different_names_close",
        "coordinate_source_disagreement",
        "coordinate_source_offset",
        "district_source_disagreement",
        "outside_tashkent_bounds",
        "dense_coordinate_cluster",
        "address_null_location_evidence",
        "district_spatial_outlier",
        "address_district_mismatch",
        "address_without_house",
    }
    coordinate_review = [
        r for r in rc if any(code in coordinate_codes for code in r["qa_warnings"])
    ]
    report = {
        "release_candidate_places": len(rc),
        "human_approval_status": "pending",
        "category_counts": dict(Counter(r["category"] for r in rc)),
        "district_counts": dict(sorted(Counter(r["district"] or "UNKNOWN" for r in rc).items())),
        "warning_events": len(warnings),
        "places_with_warnings": sum(bool(r["qa_warnings"]) for r in rc),
        "warning_counts_by_code": dict(sorted(Counter(w["code"] for w in warnings).items())),
        "coordinate_visual_review_places": len(coordinate_review),
        "field_conflict_places": sum(bool(r["field_conflicts"]) for r in rc),
        "existing_conflict_audit_events": sum(
            r["type"] == "field_conflict" for r in existing_audit
        ),
        "conflicts_resolved_from_previous_evidence": sum(
            r["type"] == "curated_field_resolution"
            and (r.get("previous_field") or {}).get("status") == "conflict"
            and r.get("new_value") is not None
            for r in existing_audit
        ),
        "blocker_places": sum(any(w["level"] == "blocker" for w in by_id[r["id"]]) for r in rc),
        "possible_duplicate_pairs": sum(w["code"] == "same_name_nearby" for w in warnings),
        "automatic_merges": 0,
    }
    write_jsonl(OUT / "places_release_candidate.jsonl", rc)
    write_jsonl(OUT / "places_qa_warnings.jsonl", warnings)
    write_jsonl(OUT / "places_coordinate_review.jsonl", coordinate_review)
    (OUT / "release_quality_report.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    _workbook(rc, warnings)
    return report


def _workbook(rows: list[dict], warnings: list[dict]) -> None:
    path = OUT / "places_review.xlsx"
    book = load_workbook(path)
    for name in ("Release candidate", "RC QA warnings"):
        if name in book:
            del book[name]
    ws = book.create_sheet("Release candidate", 0)
    columns = [
        "include",
        "source_id",
        "name",
        "category",
        "district",
        "address",
        "latitude",
        "longitude",
        "phones",
        "opening_hours",
        "website",
        "entity_confidence",
        "source_count",
        "field_conflicts",
        "qa_warning_count",
        "qa_warnings",
        "review_notes",
    ]
    ws.append(columns)
    for row in rows:
        ws.append(
            [
                (
                    ", ".join(row.get(c) or [])
                    if c in {"phones", "field_conflicts", "qa_warnings"}
                    else row.get(c)
                )
                for c in columns
            ]
        )
    issue = book.create_sheet("RC QA warnings", 1)
    issue.append(["place_ids", "code", "level", "detail"])
    for row in warnings:
        issue.append([", ".join(row["place_ids"]), row["code"], row["level"], row["detail"]])
    for sheet in (ws, issue):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="23405A")
        for column in sheet.columns:
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
                72, max(15, max(len(str(c.value or "")) for c in list(column)[:100]) + 2)
            )
    book.save(path)
    book.close()


if __name__ == "__main__":
    print(json.dumps(qa(), indent=2, ensure_ascii=False))
