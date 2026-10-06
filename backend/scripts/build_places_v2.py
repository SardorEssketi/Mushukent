"""Build a review-only Tashkent places dataset from independent V2 evidence.

No database access. Raw observations are immutable inputs; generated files live
only in data/places_v2/output. A listing is a lead, never an approved attribute.
"""

from __future__ import annotations

import json
import math
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import UTC, datetime
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data" / "places_v2"
OUT = DATA / "output"
TARGET = {"pet_store", "veterinary_clinic", "veterinary_pharmacy", "animal_shelter"}
CONTACT_FIELDS = ("phones", "website", "instagram", "telegram", "opening_hours")
DISPLAY_FIELDS = ("name", "category", "address", "coordinates", *CONTACT_FIELDS)
TASHKENT_BOUNDS = (41.15, 41.45, 69.10, 69.48)
PHONE_RE = re.compile(r"^\+998\d{9}$")


def read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    return [
        json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()
    ]


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n" for row in rows),
        encoding="utf-8",
    )


def normal(value: str) -> str:
    replacements = str.maketrans(
        {
            "а": "a",
            "б": "b",
            "в": "v",
            "г": "g",
            "д": "d",
            "е": "e",
            "ё": "e",
            "ж": "zh",
            "з": "z",
            "и": "i",
            "й": "y",
            "к": "k",
            "л": "l",
            "м": "m",
            "н": "n",
            "о": "o",
            "п": "p",
            "р": "r",
            "с": "s",
            "т": "t",
            "у": "u",
            "ф": "f",
            "х": "kh",
            "ц": "ts",
            "ч": "ch",
            "ш": "sh",
            "щ": "shch",
            "ы": "y",
            "э": "e",
            "ю": "yu",
            "я": "ya",
            "қ": "q",
            "ғ": "g",
            "ў": "o",
            "ҳ": "h",
            "ъ": "",
            "ь": "",
            "ʻ": "",
            "’": "",
            "'": "",
        }
    )
    value = unicodedata.normalize("NFKC", value).lower().translate(replacements)
    value = re.sub(
        r"\b(pro\s*\d+|oo[oо]|chp|ooo|ul|ulitsa|street|kochasi|rayon|district|tashkent|toshkent|g)\b",
        " ",
        value,
    )
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def house_numbers(address: str) -> set[str]:
    return set(re.findall(r"\b\d+[a-z]?(?:/\d+[a-z]?)?\b", normal(address)))


def address_match(a: str, b: str) -> bool:
    if not a or not b:
        return False
    na, nb = normal(a), normal(b)
    nums_a, nums_b = house_numbers(a), house_numbers(b)
    # A shared house number is insufficient when a block/mavze number differs.
    # Prefer an extra review row over merging two physical chain branches.
    if nums_a and nums_b and nums_a != nums_b:
        return False
    ta, tb = set(na.split()), set(nb.split())
    overlap = len(ta & tb) / max(1, min(len(ta), len(tb)))
    return overlap >= 0.60 or SequenceMatcher(None, na, nb).ratio() >= 0.78


def distance_meters(a: tuple[float, float], b: tuple[float, float]) -> float:
    lat1, lon1, lat2, lon2 = map(math.radians, (*a, *b))
    angular = 2 * math.asin(
        math.sqrt(
            math.sin((lat2 - lat1) / 2) ** 2
            + math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2
        )
    )
    return 6_371_000 * angular


def reviewed_firm_distances(
    link: dict, point: tuple[float, float], firm_suggestions_by_id: dict[str, dict]
) -> list[float]:
    """Require a reviewed 2GIS firm pin near an exact branch link."""
    distances = []
    for url in link["evidence_urls"]:
        match = re.search(r"2gis\.uz/tashkent/firm/(\d+)", url)
        if match and match.group(1) in firm_suggestions_by_id:
            firm_point = firm_suggestions_by_id[match.group(1)]["coordinates"]
            distance = distance_meters(point, tuple(firm_point))
            if distance > 150:
                raise ValueError(f"Candidate link is {distance:.0f} m from firm pin: {link}")
            distances.append(round(distance, 1))
    if not distances:
        raise ValueError(f"Candidate link has no reviewed 2GIS firm pin: {link}")
    return distances


def same_entity(a: dict, b: dict) -> bool:
    """Conservative merge; matching names without branch/location evidence never merge."""
    names = SequenceMatcher(None, normal(a["name"]), normal(b["name"])).ratio()
    if names < 0.89:
        return False
    coords_a = (a.get("latitude"), a.get("longitude"))
    coords_b = (b.get("latitude"), b.get("longitude"))
    close = False
    if all(value is not None for value in (*coords_a, *coords_b)):
        distance_m = distance_meters(coords_a, coords_b)
        if distance_m > 150:
            return False
        close = distance_m <= 50
    if address_match(a.get("address_hint") or "", b.get("address_hint") or ""):
        return True
    if not close:
        return False
    contacts_a = set(a.get("phones") or []) | {
        a.get("website"),
        a.get("instagram"),
        a.get("telegram"),
    }
    contacts_b = set(b.get("phones") or []) | {
        b.get("website"),
        b.get("instagram"),
        b.get("telegram"),
    }
    contacts_a.discard(None)
    contacts_b.discard(None)
    return bool(contacts_a & contacts_b)


def valid_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and "." in (parsed.hostname or "")


def field_value(item: dict, name: str):
    return item.get("fields", {}).get(name, {}).get("value")


def evidence_is_usable(field: dict) -> bool:
    return (
        field.get("status") == "verified"
        and field.get("confidence") in {"high", "medium"}
        and bool(field.get("sources"))
    )


def usable_value(item: dict, name: str):
    field = item.get("fields", {}).get(name, {})
    return field.get("value") if evidence_is_usable(field) else None


DISTRICT_HINTS = {
    "Bektemir": ("bektemir",),
    "Chilanzar": ("chilanzar", "chilonzor"),
    "Mirobad": ("mirabad", "mirobod", "oybek"),
    "Mirzo Ulugbek": ("mirzo ulugbek", "yalangach", "qorasuv", "karasu"),
    "Olmazor": ("olmazor", "almazar", "sebzar", "karakamysh"),
    "Sergeli": ("sergeli", "sirgli", "qumariq", "kumaryk"),
    "Shaykhantahur": ("shaykhantahur", "shaykhontokhur", "shaykhontokhurskiy"),
    "Uchtepa": ("uchtepa",),
    "Yakkasaray": ("yakkasaray", "yakkasarayskiy"),
    "Yangihayot": ("yangihayot", "kurgantepa", "qurgontepa"),
    "Yashnabad": ("yashnabad", "yashnobod", "aviasozlar"),
    "Yunusabad": ("yunusabad", "yunusobod"),
}


def district_hint(address: str) -> str | None:
    value = normal(address)
    hits = [
        district
        for district, terms in DISTRICT_HINTS.items()
        if any(term in value for term in terms)
    ]
    return hits[0] if len(hits) == 1 else None


def curated_house_points(rows: list[dict]) -> list[dict]:
    """Expand one-branch-per-row building evidence into field provenance."""
    items = []
    for row in rows:
        house = {"url": row["house_url"], "kind": "detailed_map_house"}
        org = {"url": row["org_url"], "kind": "map_listing"} if row.get("org_url") else None
        official = (
            {"url": row["official_url"], "kind": "official_social"}
            if row.get("official_url")
            else None
        )
        named_in_house = row.get("named_in_house", False)
        if not named_in_house and not (org or official):
            raise ValueError(f"{row['id']}: organization or official branch evidence is required")
        identity_sources = [house] if named_in_house else ([official] if official else [org])
        fields = {
            "name": {
                "value": row["name"],
                "status": "verified",
                "confidence": "high" if named_in_house or official else "medium",
                "sources": identity_sources,
            },
            "category": {
                "value": row["category"],
                "status": "verified",
                "confidence": "high" if named_in_house or official else "medium",
                "sources": identity_sources,
            },
            "address": {
                "value": row["address"],
                "status": "verified",
                "confidence": "high",
                "sources": [house] + ([org] if org else []) + ([official] if official else []),
            },
            "coordinates": {
                "value": row["coordinates"],
                "status": "verified",
                "confidence": "high",
                "sources": [house],
            },
        }
        if row.get("phones"):
            if not org:
                raise ValueError(f"{row['id']}: phone needs a branch organization URL")
            fields["phones"] = {
                "value": row["phones"],
                "status": "verified",
                "confidence": "medium",
                "sources": [org],
            }
        items.append(
            {
                "id": row["id"],
                "name": row["name"],
                "category": row["category"],
                "district": row["district"],
                "branch_name": row.get("branch_name"),
                "secondary_categories": row.get("secondary_categories", []),
                "candidate_observation_ids": row.get("candidate_observation_ids", []),
                "fields": fields,
            }
        )
    return items


def curated_org_points(rows: list[dict]) -> list[dict]:
    """Turn individually reviewed public organization pins into sourced places."""
    items = []
    for row in rows:
        org = {"url": row["org_url"], "kind": row.get("source_kind", "detailed_map_organization")}
        official = (
            {"url": row["official_url"], "kind": "official_website"}
            if row.get("official_url")
            else None
        )
        fields = {
            key: {"value": value, "status": "verified", "confidence": "high", "sources": [org]}
            for key, value in (
                ("name", row["name"]),
                ("category", row["category"]),
                ("coordinates", row["coordinates"]),
                ("address", row.get("address")),
            )
            if value is not None
        }
        if official:
            for key in ("name", "category", "address"):
                if key in fields:
                    fields[key]["sources"].append(official)
        if row.get("address_conflict"):
            fields["address"] = {
                "value": None,
                "status": "conflict",
                "confidence": "conflict",
                "sources": [org]
                + [
                    {"url": source["url"], "kind": source["kind"]}
                    for source in row["address_conflict"]
                ],
                "alternatives": row["address_conflict"],
            }
            fields["location_evidence"] = {
                "value": "organization-specific map pin",
                "status": "verified",
                "confidence": "high",
                "sources": [org],
            }
        elif row.get("address_unavailable"):
            fields["location_evidence"] = {
                "value": "firm-specific route point with public district listing",
                "status": "verified",
                "confidence": "high",
                "sources": [org],
            }
        if row.get("phones"):
            fields["phones"] = {
                "value": row["phones"],
                "status": "verified",
                "confidence": "medium",
                "sources": [org],
            }
        items.append(
            {
                "id": row["id"],
                "name": row["name"],
                "category": row["category"],
                "district": row["district"],
                "branch_name": row.get("branch_name"),
                "secondary_categories": row.get("secondary_categories", []),
                "candidate_observation_ids": row.get("candidate_observation_ids", []),
                "fields": fields,
            }
        )
    return items


def reviewed_2gis_chain_points(suggestions: list[dict], selections: list[dict]) -> list[dict]:
    """Require explicit branch and point selections before map promotion."""
    if len({row["firm_id"] for row in selections}) != len(selections):
        raise ValueError("A 2GIS firm was selected as more than one physical place")
    by_firm = {row["firm_id"]: row for row in suggestions}
    rows = []
    for selected in selections:
        row = by_firm[selected["firm_id"]]
        if (
            row["coordinates"] != selected["expected_coordinates"]
            or row["address"] != selected["expected_address"]
        ):
            raise ValueError(f"2GIS branch evidence changed: {selected['firm_id']}")
        rows.append(
            {
                "id": selected["id"],
                "name": row["name"],
                "category": row["category"],
                "district": selected["district"],
                "branch_name": selected["branch_name"],
                "secondary_categories": selected.get("secondary_categories", []),
                "address": "Tashkent, " + row["address"],
                "coordinates": row["coordinates"],
                "org_url": row["source_url"],
                "source_kind": "detailed_2gis_organization",
                "official_url": selected.get("official_url"),
                "address_conflict": selected.get("address_conflict"),
                "candidate_observation_ids": ["2gis_" + row["firm_id"]]
                + selected.get("additional_observation_ids", []),
            }
        )
    return curated_org_points(rows)


def build() -> dict:
    OUT.mkdir(parents=True, exist_ok=True)
    two_gis_suggestions = sum(
        (
            read_jsonl(DATA / f"2gis_{brand}_firm_suggestions.jsonl")
            for brand in (
                "zooplaneta",
                "animalplanet",
                "arcazoo",
                "felix",
                "neko",
                "district_clinic",
                "district_petstore",
                "district_pharmacy",
            )
        ),
        [],
    )
    if len({row["firm_id"] for row in two_gis_suggestions}) != len(two_gis_suggestions):
        raise ValueError(
            "Duplicate 2GIS firm IDs across discovery sources need explicit reconciliation"
        )
    all_observations = sum(
        (
            read_jsonl(DATA / filename)
            for filename in (
                "source_observations.jsonl",
                "map_observations.jsonl",
                "yandex_pet_pages_2_5.jsonl",
                "yandex_district_observations.jsonl",
                "yandex_district_extra.jsonl",
                "yandex_vet_pharmacy_page_2.jsonl",
                "yandex_vet_pharmacy_pages_3plus.jsonl",
                "yandex_vet_clinic_pages_2_3.jsonl",
                "yandex_pet_goods_page_6.jsonl",
                "official_chain_observations.jsonl",
                "yandex_pet_goods_pages_7plus.jsonl",
                "yandex_veterinary_drugs_observations.jsonl",
                "felix_chain_and_feruza_observations.jsonl",
                "zooplaneta_branch_observations.jsonl",
                "chain_house_observations.jsonl",
                "org_targeted_discovery.jsonl",
                "zoom_telegram_observations.jsonl",
                "shelter_and_other_observations.jsonl",
            )
        ),
        [],
    )
    all_observations += [
        {
            "observation_id": "2gis_" + row["firm_id"],
            "name": row["name"],
            "address_hint": "Tashkent, " + row["address"],
            "category_hint": row["category"],
            "coordinates_hint": row["coordinates"],
            "source_url": row["source_url"],
            "source_kind": "2gis_branch_firm",
            "discovery_status": "unverified",
        }
        for row in two_gis_suggestions
    ]
    curated = json.loads((DATA / "curated_evidence.json").read_text(encoding="utf-8"))
    curated += read_jsonl(DATA / "additional_verified_evidence.jsonl")
    curated += read_jsonl(DATA / "felix_feruza_verified_evidence.jsonl")
    curated += read_jsonl(DATA / "felix_mirobod_verified_evidence.jsonl")
    curated += read_jsonl(DATA / "zooplaneta_verified_evidence.jsonl")
    curated += curated_house_points(read_jsonl(DATA / "house_verified_points.jsonl"))
    curated += curated_house_points(read_jsonl(DATA / "yandex_house_points_reviewed.jsonl"))
    curated += curated_org_points(read_jsonl(DATA / "org_verified_points.jsonl"))
    curated += curated_org_points(read_jsonl(DATA / "yandex_org_points_reviewed.jsonl"))
    curated += curated_org_points(read_jsonl(DATA / "2gis_unaddressed_clinics_verified.jsonl"))
    curated += reviewed_2gis_chain_points(
        two_gis_suggestions,
        sum(
            (
                read_jsonl(DATA / f"2gis_{brand}_verified_selections.jsonl")
                for brand in (
                    "zooplaneta",
                    "animalplanet",
                    "arcazoo",
                    "felix",
                    "neko",
                    "district_clinic",
                    "district_petstore",
                    "district_pharmacy",
                )
            ),
            [],
        ),
    )
    curated_by_id = {item["id"]: item for item in curated}
    if len(curated_by_id) != len(curated):
        raise ValueError("Duplicate curated physical-place IDs")
    resolution_audit = []
    for resolution in (
        read_jsonl(DATA / "curated_review_resolutions.jsonl")
        + read_jsonl(DATA / "curated_review_resolutions_additional.jsonl")
        + read_jsonl(DATA / "curated_review_resolutions_links.jsonl")
        + read_jsonl(DATA / "curated_review_resolutions_release.jsonl")
    ):
        item = curated_by_id[resolution["place_id"]]
        item.pop("review_status", None)
        item.pop("review_reason", None)
        item["review_note"] = resolution["reason"]
        item.setdefault("candidate_observation_ids", []).extend(
            resolution.get("candidate_observation_ids", [])
        )
        for field_name, sources in resolution.get("append_sources", {}).items():
            item["fields"][field_name].setdefault("sources", []).extend(sources)
        for field_name, field in resolution["fields"].items():
            previous = item["fields"].get(field_name)
            resolution_audit.append(
                {
                    "type": "curated_field_resolution",
                    "place_id": item["id"],
                    "field": field_name,
                    "old_value": previous.get("value") if previous else None,
                    "previous_field": previous,
                    "new_value": field.get("value"),
                    "sources": field.get("sources", []),
                    "reason": resolution["reason"],
                }
            )
            item["fields"][field_name] = field
    observations_by_id = {row["observation_id"]: row for row in all_observations}
    firm_suggestions_by_id = {row["firm_id"]: row for row in two_gis_suggestions}
    for link in read_jsonl(DATA / "candidate_exact_branch_links.jsonl"):
        if link["candidate_id"] != "v2_candidate_" + link["observation_id"]:
            raise ValueError(f"Candidate link ID mismatch: {link}")
        observation = observations_by_id[link["observation_id"]]
        item = curated_by_id[link["place_id"]]
        if observation["category_hint"] not in {
            item["category"],
            *item.get("secondary_categories", []),
        }:
            raise ValueError(f"Candidate link category mismatch: {link}")
        point = item["fields"]["coordinates"]["value"]
        firm_distances = reviewed_firm_distances(link, tuple(point), firm_suggestions_by_id)
        item.setdefault("candidate_observation_ids", []).append(link["observation_id"])
        resolution_audit.append(
            {
                "type": "candidate_exact_branch_link",
                **link,
                "firm_point_distances_m": firm_distances,
            }
        )
    for addition in read_jsonl(DATA / "field_enrichment_additions.jsonl"):
        item = curated_by_id[addition["place_id"]]
        old = item["fields"].get(addition["field"])
        if old and (old.get("value") is not None or old.get("status") == "conflict"):
            raise ValueError(
                "Field enrichment would overwrite reviewed evidence: "
                f"{addition['place_id']} {addition['field']}"
            )
        item["fields"][addition["field"]] = {
            key: addition[key] for key in ("value", "status", "confidence", "sources")
        }
    for extra in read_jsonl(DATA / "supplementary_category_evidence.jsonl") + read_jsonl(
        DATA / "supplementary_category_evidence_yandex.jsonl"
    ):
        if extra["secondary_category"] not in TARGET:
            raise ValueError(f"Unsupported supplemental category: {extra}")
        item = curated_by_id[extra["place_id"]]
        categories = item.setdefault("secondary_categories", [])
        if extra["secondary_category"] not in categories:
            categories.append(extra["secondary_category"])
        item["fields"]["secondary_category:" + extra["secondary_category"]] = {
            "value": extra["secondary_category"],
            "status": "verified",
            "confidence": "medium",
            "sources": [
                {
                    "url": extra["source_url"],
                    "kind": extra.get("source_kind", "detailed_2gis_organization"),
                }
            ],
        }
    research_decisions = read_jsonl(DATA / "candidate_research_decisions.jsonl")
    decisions_by_candidate = {row["candidate_id"]: row for row in research_decisions}
    probes = read_jsonl(DATA / "yandex_lead_probes.jsonl") + read_jsonl(
        DATA / "yandex_lead_probes_address.jsonl"
    )
    probe_by_candidate = {row["candidate_id"]: row for row in probes}
    stamp = datetime.now(UTC).isoformat()
    audit: list[dict] = list(research_decisions) + resolution_audit
    source_rows: list[dict] = []
    for probe in probes:
        audit.append(
            {
                "type": "public_search_probe",
                "candidate_id": probe["candidate_id"],
                "query_url": probe["search_url"],
                "result_url": probe.get("organization_url"),
                "probe_status": probe["probe_status"],
                "decision": "suggestion only; no entity match or attribute promotion",
            }
        )
        if probe.get("organization_url"):
            source_rows.append(
                {
                    "record_id": probe["candidate_id"],
                    "field": "search_result_suggestion",
                    "value": {
                        "title": probe.get("result_title"),
                        "address": probe.get("result_address"),
                        "coordinates": probe.get("result_coordinates"),
                    },
                    "source_url": probe["organization_url"],
                    "source_kind": "public_search_probe",
                    "status": "unverified_suggestion",
                    "confidence": "low",
                    "checked_at": probe["checked_at"],
                }
            )
    for item in curated:
        for url in item.get("supporting_sources", []):
            for key in ("name", "category", "address"):
                item["fields"][key]["sources"].append({"url": url, "kind": "map_listing"})
        for key, field in item["fields"].items():
            if field.get("confidence") != "high":
                continue
            providers = {
                urlparse(src["url"]).hostname.removeprefix("www.")
                for src in field.get("sources", [])
            }
            official = any(
                src["kind"] in {"official_website", "official_social"}
                for src in field.get("sources", [])
            )
            detailed_house = key in {"name", "category", "address", "coordinates"} and any(
                src["kind"] == "detailed_map_house" for src in field.get("sources", [])
            )
            if not official and not detailed_house and len(providers) < 2:
                field["confidence"] = "medium"
                audit.append(
                    {
                        "type": "field_confidence_downgraded",
                        "place_id": item["id"],
                        "field": key,
                        "reason": "one independent nonofficial provider supports this field",
                    }
                )
    exclusions = {
        "dir_0020": "directory navigation header, not a business",
        "dir_0037": "agricultural company; pet-facing pharmacy not established",
        "dir_0060": "grooming service, outside four target categories",
        "dir_0064": "pet beauty salon, outside four target categories",
        "dir_0092": "pet beauty salon, outside four target categories",
        "shelter_003": "documented in Piskent district outside Tashkent city",
        "shelter_004": "animal welfare organization's legal address is not a public shelter",
    }
    observations = []
    other_categories = []
    for obs in all_observations:
        reason = exclusions.get(obs["observation_id"])
        if reason or obs["category_hint"] not in TARGET:
            other_categories.append(
                {**obs, "exclusion_reason": reason or "outside target categories"}
            )
            audit.append(
                {
                    "type": "discovery_exclusion",
                    "observation_id": obs["observation_id"],
                    "reason": reason or "outside target categories",
                }
            )
        else:
            observations.append(obs)

    # Each observation is tied to one name, address hint, category hint and URL.
    # Merge only same-name + same-address leads; divergent branch addresses stay apart.
    groups: list[list[dict]] = []
    for obs in observations:
        if obs["category_hint"] not in TARGET:
            groups.append([obs])
            continue
        match = next((group for group in groups if same_entity(group[0], obs)), None)
        if match is None:
            groups.append([obs])
        else:
            match.append(obs)
            audit.append(
                {
                    "type": "duplicate_observation",
                    "observation_id": obs["observation_id"],
                    "retained_observation_id": match[0]["observation_id"],
                    "reason": (
                        "same normalized name and compatible house/address; "
                        "physical branches preserved"
                    ),
                }
            )

    candidates: list[dict] = []
    for group in groups:
        first = group[0]
        candidate = {
            "candidate_id": "v2_candidate_" + first["observation_id"],
            "name_hint": first["name"],
            "address_hint": first.get("address_hint"),
            "category_hint": first["category_hint"],
            "district_hint": district_hint(first.get("address_hint") or ""),
            "observation_ids": [x["observation_id"] for x in group],
            "source_urls": sorted({x["source_url"] for x in group}),
            "status": "discovery_only",
            "verified_place_id": None,
        }
        candidates.append(candidate)
        candidate["source_count"] = len(candidate["source_urls"])
        for obs in group:
            source_rows.append(
                {
                    "record_id": candidate["candidate_id"],
                    "field": "discovery_hint",
                    "value": {
                        "name": obs["name"],
                        "address": obs.get("address_hint"),
                        "category": obs["category_hint"],
                    },
                    "source_url": obs["source_url"],
                    "source_kind": obs["source_kind"],
                    "status": "discovery_only",
                    "confidence": "low",
                    "checked_at": stamp,
                }
            )
    for index, left in enumerate(candidates):
        for right in candidates[index + 1 :]:
            similarity = SequenceMatcher(
                None, normal(left["name_hint"]), normal(right["name_hint"])
            ).ratio()
            if 0.78 <= similarity < 0.89 and address_match(
                left.get("address_hint") or "", right.get("address_hint") or ""
            ):
                audit.append(
                    {
                        "type": "possible_duplicate_candidate",
                        "candidate_ids": [left["candidate_id"], right["candidate_id"]],
                        "reason": (
                            "similar names and address; identity is not certain enough to merge"
                        ),
                    }
                )

    verified: list[dict] = []
    needs_review: list[dict] = []
    for item in curated:
        fields = item["fields"]
        if item.get("review_note"):
            audit.append(
                {
                    "type": "curated_review_note",
                    "place_id": item["id"],
                    "reason": item["review_note"],
                }
            )
        for note in item.get("review_note_sources", []):
            source_rows.append(
                {
                    "record_id": item["id"],
                    "field": note["field"],
                    "value": note["value"],
                    "source_url": note["source_url"],
                    "source_kind": "review_note",
                    "status": note["status"],
                    "confidence": "low",
                    "checked_at": stamp,
                }
            )
        place = {
            "id": item["id"],
            "name": usable_value(item, "name"),
            "branch_name": item.get("branch_name"),
            "alternative_names": item.get("alternative_names", []),
            "category": usable_value(item, "category"),
            "secondary_categories": item.get("secondary_categories", []),
            "latitude": None,
            "longitude": None,
            "address": usable_value(item, "address"),
            "district": item.get("district"),
            "phones": usable_value(item, "phones") or [],
            "website": usable_value(item, "website"),
            "instagram": usable_value(item, "instagram"),
            "telegram": usable_value(item, "telegram"),
            "opening_hours": usable_value(item, "opening_hours"),
            "description": usable_value(item, "description"),
            "verified_at": stamp,
            "source_urls": sorted(
                {s["url"] for f in fields.values() for s in f.get("sources", [])}
            ),
            "field_confidence": {key: field.get("confidence") for key, field in fields.items()},
            "entity_confidence": "high",
            "source_count": len({s["url"] for f in fields.values() for s in f.get("sources", [])}),
            "review_status": item.get("review_status", "verified"),
        }
        coords = usable_value(item, "coordinates")
        if isinstance(coords, list) and len(coords) == 2:
            place["latitude"], place["longitude"] = coords
        place["missing_phone"] = not bool(place["phones"])
        place["missing_hours"] = not bool(place["opening_hours"])
        place["field_conflicts"] = sorted(
            key for key, field in fields.items() if field.get("status") == "conflict"
        )

        for key, field in fields.items():
            for src in field.get("sources", []):
                source_rows.append(
                    {
                        "record_id": item["id"],
                        "field": key,
                        "value": field.get("value"),
                        "source_url": src["url"],
                        "source_kind": src["kind"],
                        "branch_scope": src.get("branch_scope"),
                        "branch_label": src.get("branch_label"),
                        "status": field.get("status"),
                        "confidence": field.get("confidence"),
                        "checked_at": stamp,
                    }
                )
            if field.get("status") == "conflict":
                audit.append(
                    {
                        "type": "field_conflict",
                        "place_id": item["id"],
                        "field": key,
                        "alternatives": field.get("alternatives", []),
                        "decision": "set_null",
                    }
                )
                for alternative in field.get("alternatives", []):
                    if isinstance(alternative, dict) and alternative.get("url"):
                        source_rows.append(
                            {
                                "record_id": item["id"],
                                "field": key,
                                "value": alternative["value"],
                                "source_url": alternative["url"],
                                "source_kind": alternative.get("kind", "map_listing"),
                                "status": "conflict_alternative",
                                "confidence": "conflict",
                                "checked_at": stamp,
                            }
                        )
            for alternate in field.get("unconfirmed_alternatives", []):
                source_rows.append(
                    {
                        "record_id": item["id"],
                        "field": key,
                        "value": alternate["value"],
                        "source_url": alternate["source_url"],
                        "source_kind": "map_listing",
                        "status": "unconfirmed_alternative",
                        "confidence": "low",
                        "checked_at": stamp,
                    }
                )
                audit.append(
                    {
                        "type": "unconfirmed_alternative",
                        "place_id": item["id"],
                        "field": key,
                        "value": alternate["value"],
                        "reason": "not attributed to this branch by official source",
                    }
                )
        for derived_field, source_field in (("district", "address"), ("branch_name", "name")):
            if not item.get(derived_field):
                continue
            source = fields.get(source_field) or fields.get("location_evidence", {})
            for src in source.get("sources", []):
                source_rows.append(
                    {
                        "record_id": item["id"],
                        "field": derived_field,
                        "value": item[derived_field],
                        "source_url": src["url"],
                        "source_kind": src["kind"],
                        "status": "derived_for_review",
                        "confidence": "medium",
                        "checked_at": stamp,
                    }
                )

        reasons = []
        if item.get("review_reason"):
            reasons.append(item["review_reason"])
        for field in fields.values():
            for src in field.get("sources", []):
                scope = src.get("branch_scope")
                if scope and normal(scope) != normal(item.get("branch_name") or ""):
                    reasons.append(
                        f"source branch scope {scope!r} differs from this physical branch"
                    )
        if not place["name"] or place["category"] not in TARGET:
            reasons.append("missing verified identity or target category")
        has_location_evidence = evidence_is_usable(fields.get("location_evidence", {}))
        if (
            (not place["address"] and not has_location_evidence)
            or place["latitude"] is None
            or place["longitude"] is None
        ):
            reasons.append("missing verified physical address or coordinates")
        if place["latitude"] is not None:
            lat, lon = place["latitude"], place["longitude"]
            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                reasons.append("invalid geographic coordinates")
            if not (
                TASHKENT_BOUNDS[0] <= lat <= TASHKENT_BOUNDS[1]
                and TASHKENT_BOUNDS[2] <= lon <= TASHKENT_BOUNDS[3]
            ):
                reasons.append("outside conservative Tashkent review bounds")
        for phone in place["phones"]:
            if not PHONE_RE.fullmatch(phone):
                reasons.append(f"invalid Uzbekistan phone: {phone}")
        for key in ("website", "instagram", "telegram"):
            if place[key] and not valid_url(place[key]):
                reasons.append(f"invalid {key} URL")
        if any(not valid_url(src) for src in place["source_urls"]):
            reasons.append("invalid provenance URL")
        for key in ("name", "category", "coordinates"):
            if not evidence_is_usable(fields.get(key, {})):
                reasons.append(f"{key} has no usable source evidence")
        if not evidence_is_usable(fields.get("address", {})) and not has_location_evidence:
            reasons.append("address has no usable source evidence")

        if reasons:
            place["review_status"] = "needs_review"
            place["entity_confidence"] = "needs_review"
            place["review_reasons"] = sorted(set(reasons))
            needs_review.append(place)
            audit.append(
                {
                    "type": "curated_record_held",
                    "place_id": item["id"],
                    "reasons": place["review_reasons"],
                }
            )
        else:
            place["review_status"] = "verified"
            verified.append(place)
            audit.append(
                {
                    "type": "curated_record_approved_for_review",
                    "place_id": item["id"],
                    "reason": (
                        "source-backed identity/category/address/coordinates passed validation"
                    ),
                }
            )
            conflicted = sorted(
                key for key, field in fields.items() if field.get("status") == "conflict"
            )
            if conflicted:
                needs_review.append(
                    {
                        **place,
                        "review_status": "field_conflict",
                        "review_reasons": [
                            f"Resolve {', '.join(conflicted)} before publishing those fields"
                        ],
                        "conflicted_fields": conflicted,
                    }
                )

        # Link leads to curated places only when name and house/address agree.
        linked = False
        for candidate in candidates:
            candidate_as_obs = {
                "name": candidate["name_hint"],
                "address_hint": candidate["address_hint"],
            }
            curated_as_obs = {"name": item["name"], "address_hint": place["address"]}
            if set(candidate["observation_ids"]) & set(
                item.get("candidate_observation_ids", [])
            ) or same_entity(candidate_as_obs, curated_as_obs):
                candidate["verified_place_id"] = item["id"]
                candidate["status"] = place["review_status"]
                candidate["district_hint"] = item.get("district")
                linked = True
        if not linked:
            candidates.append(
                {
                    "candidate_id": "v2_candidate_" + item["id"],
                    "name_hint": item["name"],
                    "address_hint": place["address"],
                    "category_hint": item["category"],
                    "district_hint": item.get("district"),
                    "observation_ids": [],
                    "source_urls": place["source_urls"],
                    "status": place["review_status"],
                    "verified_place_id": item["id"],
                    "source_count": len(place["source_urls"]),
                }
            )
            audit.append(
                {
                    "type": "candidate_from_targeted_verification",
                    "place_id": item["id"],
                    "reason": (
                        "curated independently from a source page absent from raw "
                        "directory harvest"
                    ),
                }
            )

    for candidate in candidates:
        if not candidate["verified_place_id"]:
            decision = decisions_by_candidate.get(candidate["candidate_id"])
            if decision and decision["type"] in {
                "candidate_outside_city",
                "candidate_outside_product_scope",
                "candidate_closed",
            }:
                candidate["status"] = "excluded_from_map"
            elif decision and decision["type"] in {
                "candidate_possible_alias",
                "candidate_needs_review",
            }:
                candidate["status"] = "needs_review"
            needs_review.append(
                {
                    "id": candidate["candidate_id"],
                    "name_hint": candidate["name_hint"],
                    "category_hint": candidate["category_hint"],
                    "address_hint": candidate["address_hint"],
                    "district_hint": candidate.get("district_hint"),
                    "latitude": None,
                    "longitude": None,
                    "review_status": candidate["status"],
                    "review_reasons": (
                        [decision["reason"]]
                        if decision
                        else [
                            "identity, current operation and precise coordinates not "
                            "independently verified"
                        ]
                    ),
                    "source_urls": candidate["source_urls"],
                    "source_count": candidate["source_count"],
                    "entity_confidence": "unknown",
                    "missing_phone": True,
                    "missing_hours": True,
                    "field_conflicts": [],
                }
            )

    # One row per lead makes the remaining verification work explicit and filterable.
    observation_by_id = {o["observation_id"]: o for o in observations}
    triage = []
    for candidate in candidates:
        kinds = sorted(
            {observation_by_id[oid]["source_kind"] for oid in candidate["observation_ids"]}
        )
        if candidate["status"] == "verified":
            next_action = "review_field_evidence"
        elif candidate["status"] == "needs_review":
            next_action = "resolve_curated_entity_or_address_conflict"
        elif candidate["status"] == "excluded_from_map":
            next_action = "review_exclusion_only_if_new_evidence_appears"
        elif candidate["category_hint"] == "animal_shelter":
            next_action = "confirm_explicit_public_visitor_location_before_map"
        elif candidate["candidate_id"] in probe_by_candidate:
            next_action = "compare_public_search_result_to_exact_branch_then_curate_point"
        elif any(
            kind in {"yandex_search", "yandex_district", "2gis_search", "targeted_web_search"}
            for kind in kinds
        ):
            next_action = "open_detailed_map_listing_and_confirm_exact_branch_pin"
        else:
            next_action = "find_current_map_or_official_business_listing_and_branch_pin"
        triage.append(
            {
                "candidate_id": candidate["candidate_id"],
                "name_hint": candidate["name_hint"],
                "category_hint": candidate["category_hint"],
                "district_hint": candidate.get("district_hint"),
                "status": candidate["status"],
                "source_kinds": kinds,
                "research_status": (
                    "source_backed_point"
                    if candidate["status"] == "verified"
                    else (
                        "ambiguous_candidate_held"
                        if candidate["status"] == "needs_review"
                        else (
                            "excluded_after_research"
                            if candidate["status"] == "excluded_from_map"
                            else (
                                "attempted_location_held"
                                if candidate["candidate_id"] in decisions_by_candidate
                                else (
                                    "public_search_probed_unresolved"
                                    if candidate["candidate_id"] in probe_by_candidate
                                    else "not_yet_checked_individually"
                                )
                            )
                        )
                    )
                ),
                "source_count": candidate["source_count"],
                "next_action": next_action,
                "note": decisions_by_candidate.get(candidate["candidate_id"], {}).get("reason")
                or (
                    "Core identity and location are source-backed; review optional fields "
                    "and branch attribution."
                    if candidate["status"] == "verified"
                    else (
                        "A public search was attempted; any result is a suggestion only "
                        "until identity, branch, category and pin are matched."
                        if candidate["candidate_id"] in probe_by_candidate
                        else (
                            "Discovery evidence is insufficient for an importable point "
                            "until a sourced branch coordinate is recorded."
                        )
                    )
                ),
            }
        )

    # Shared building coordinates are normal. Shared contacts across unrelated names are not.
    hold_ids: set[str] = set()
    for key in ("phones", "website", "instagram", "telegram"):
        owner: dict[str, list[str]] = defaultdict(list)
        for place in verified:
            values = (
                place[key] if isinstance(place[key], list) else ([place[key]] if place[key] else [])
            )
            for value in values:
                owner[str(value).lower()].append(place["id"])
        for value, ids in owner.items():
            names = {
                normal(next(p["name"] for p in verified if p["id"] == place_id)) for place_id in ids
            }
            if len(ids) > 1 and len(names) > 1:
                audit.append(
                    {
                        "type": "shared_contact_review",
                        "field": key,
                        "value": value,
                        "place_ids": ids,
                        "reason": "verify branch-level attribution before import",
                    }
                )
                # A shared company contact may be legitimate across physical branches.
                # Review the field attribution without rejecting source-backed locations.
    for i, a in enumerate(verified):
        for b in verified[i + 1 :]:
            if (a["latitude"], a["longitude"]) == (b["latitude"], b["longitude"]):
                audit.append(
                    {
                        "type": "shared_building_coordinates",
                        "place_ids": [a["id"], b["id"]],
                        "reason": "same building; distinct named businesses remain separate",
                    }
                )
                if normal(a["name"]) == normal(b["name"]):
                    hold_ids.update((a["id"], b["id"]))
                    audit.append(
                        {
                            "type": "duplicate_curated_identity",
                            "place_ids": [a["id"], b["id"]],
                            "reason": "same normalized name and coordinate; both held for review",
                        }
                    )
            if normal(a["name"]) == normal(b["name"]) and a["address"] != b["address"]:
                audit.append(
                    {
                        "type": "chain_branches_preserved",
                        "place_ids": [a["id"], b["id"]],
                        "reason": "same brand, distinct physical addresses",
                    }
                )
    if hold_ids:
        for place in verified:
            if place["id"] in hold_ids:
                needs_review.append(
                    {
                        **place,
                        "review_status": "needs_review",
                        "review_reasons": [
                            "duplicate identity or reused contact across unrelated businesses"
                        ],
                    }
                )
        verified = [place for place in verified if place["id"] not in hold_ids]
        for candidate in candidates:
            if candidate["verified_place_id"] in hold_ids:
                candidate["status"] = "needs_review"

    # Search-log counts are observations by source page, not an invented claim of completeness.
    source_stats: dict[str, dict] = {}
    seen_candidates: set[str] = set()
    for candidate in candidates:
        for source_url in candidate["source_urls"]:
            stat = source_stats.setdefault(
                source_url,
                {
                    "source_url": source_url,
                    "observations": 0,
                    "candidate_ids": set(),
                    "first_seen_unique": 0,
                },
            )
            stat["candidate_ids"].add(candidate["candidate_id"])
    for item in curated:
        for field in item["fields"].values():
            for src in field.get("sources", []):
                source_stats.setdefault(
                    src["url"],
                    {
                        "source_url": src["url"],
                        "observations": 0,
                        "candidate_ids": set(),
                        "first_seen_unique": 0,
                        "source_kind": src["kind"],
                    },
                )
    for obs in observations:
        source_stats[obs["source_url"]]["observations"] += 1
        candidate = next(c for c in candidates if obs["observation_id"] in c["observation_ids"])
        if candidate["candidate_id"] not in seen_candidates:
            source_stats[obs["source_url"]]["first_seen_unique"] += 1
            seen_candidates.add(candidate["candidate_id"])
    search_log = []
    for url, stat in source_stats.items():
        matching_obs = next((o for o in observations if o["source_url"] == url), None)
        search_log.append(
            {
                "source_url": url,
                "source_kind": (
                    matching_obs["source_kind"]
                    if matching_obs
                    else stat.get("source_kind", "targeted_verification")
                ),
                "query_or_page": urlparse(url).path,
                "observations": stat["observations"],
                "candidate_locations_on_page": len(stat["candidate_ids"]),
                "first_seen_unique_candidates": stat["first_seen_unique"],
                "checked_at": stamp,
                "note": "Page-level result, not proof of coverage saturation",
            }
        )
    for query in read_jsonl(DATA / "search_queries.jsonl"):
        search_log.append(
            {
                "query_id": query["query_id"],
                "query": query["query"],
                "scope": query["scope"],
                "status": query["status"],
                "note": query["note"],
            }
        )
    for query in read_jsonl(DATA / "executed_queries.jsonl"):
        search_log.append(query)
    for query in read_jsonl(DATA / "continuation_queries.jsonl"):
        search_log.append(query)
    for query in read_jsonl(DATA / "district_sweep_queries.jsonl"):
        search_log.append(query)
    pagination_yields = sum(
        (
            read_jsonl(DATA / filename)
            for filename in (
                "yandex_page_yields.jsonl",
                "yandex_vet_clinic_page_yields.jsonl",
                "yandex_vet_pharmacy_page_yields.jsonl",
            )
        ),
        [],
    )
    for index, page in enumerate(pagination_yields):
        search_log.append(
            {
                "query_id": f"public_page_{index + 1}",
                "query": page["source_url"],
                "scope": "Tashkent public map pagination",
                "status": "executed",
                "new_candidate_yield": page["first_seen_name_address_pairs"],
                "target_observations": page["target_observations"],
                "checked_at": page["checked_at"],
                "result_note": page["note"],
            }
        )
    two_gis_yields = sum(
        (
            read_jsonl(DATA / f"2gis_{brand}_branch_yields.jsonl")
            for brand in ("zooplaneta", "animalplanet", "arcazoo", "felix", "neko")
        ),
        [],
    )
    for index, page in enumerate(two_gis_yields):
        search_log.append(
            {
                "query_id": f"2gis_branch_page_{index + 1}",
                "query": page["source_url"],
                "scope": "2GIS chain branch enumeration",
                "status": "executed",
                "first_seen_firm_links_within_sweep": page["first_seen_firm_links"],
                "result_note": (
                    "Firm-link yield within this sweep; repeated sweep URLs are not "
                    "additive and links are not necessarily new physical places."
                ),
                "checked_at": page["checked_at"],
            }
        )
    district_search_yields = read_jsonl(DATA / "2gis_district_clinic_search_yields.jsonl")
    for index, page in enumerate(district_search_yields):
        search_log.append(
            {
                "query_id": f"2gis_district_clinic_{index + 1}",
                "query": page["source_url"],
                "scope": page["district"],
                "status": "executed",
                "result_links": page["result_links"],
                "pet_relevant_result_links": page["pet_relevant_result_links"],
                "new_verified_physical_places": page["new_verified_physical_places"],
                "result_note": page["note"],
                "checked_at": page["checked_at"],
            }
        )
    petstore_search_yields = read_jsonl(DATA / "2gis_district_petstore_search_yields.jsonl")
    for index, page in enumerate(petstore_search_yields):
        search_log.append(
            {
                "query_id": f"2gis_district_petstore_{index + 1}",
                "query": page["source_url"],
                "scope": page["district"],
                "status": "executed",
                "result_links": page["result_links"],
                "new_verified_physical_places": page["new_verified_physical_places"],
                "result_note": page["note"],
                "checked_at": page["checked_at"],
            }
        )
    petstore_selections = {
        row["firm_id"]
        for row in read_jsonl(DATA / "2gis_district_petstore_verified_selections.jsonl")
    }
    petstore_pagination_checks = read_jsonl(
        DATA / "2gis_district_petstores_pagination_checks.jsonl"
    )
    for index, page in enumerate(petstore_pagination_checks):
        page_file = page["search_file"].replace(
            "_petstores_probe.html", f"_page{page['page']}_petstores_probe.html"
        )
        accepted = sum(
            row["firm_id"] in petstore_selections and row.get("discovery_search_file") == page_file
            for row in two_gis_suggestions
        )
        search_log.append(
            {
                "query_id": f"2gis_district_petstore_page_{index + 1}",
                "query": page["source_url"],
                "scope": page["search_file"],
                "status": page["status"],
                "result_links": page.get("result_links"),
                "new_firm_links_vs_prior_pages": page.get("new_firm_links_vs_prior_pages"),
                "new_verified_physical_places": accepted,
                "result_note": (
                    "Firm IDs are discovery only; accepted locations have reviewed "
                    "firm-specific pins."
                    if page["status"] == "executed"
                    else f"HTTP {page.get('http_status')}; yield unknown."
                ),
                "checked_at": page["checked_at"],
            }
        )
    pharmacy_selections = read_jsonl(DATA / "2gis_district_pharmacy_verified_selections.jsonl")
    pharmacy_selection_by_district = Counter(row["district"] for row in pharmacy_selections)
    pharmacy_search_checks = read_jsonl(DATA / "2gis_district_pharmacy_search_checks.jsonl")
    for index, page in enumerate(pharmacy_search_checks):
        search_log.append(
            {
                "query_id": f"2gis_district_pharmacy_{index + 1}",
                "query": page["source_url"],
                "scope": page["district"],
                "status": page["status"],
                "result_links": page.get("result_links"),
                "new_verified_physical_places": pharmacy_selection_by_district[page["district"]],
                "result_note": (
                    "Visible first page; firm-specific category, address and pin "
                    "reviewed before selection."
                    if page["status"] == "executed"
                    else f"HTTP {page.get('http_status')}; yield unknown."
                ),
                "checked_at": page["checked_at"],
            }
        )
    pharmacy_pagination_checks = read_jsonl(
        DATA / "2gis_district_pharmacies_pagination_checks.jsonl"
    )
    pharmacy_firms = {row["firm_id"]: row for row in two_gis_suggestions}
    selected_pharmacy_firms = {row["firm_id"] for row in pharmacy_selections}
    for index, page in enumerate(pharmacy_pagination_checks):
        page_file = page["search_file"].replace(
            "_pharmacies_probe.html", f"_page{page['page']}_pharmacies_probe.html"
        )
        accepted = sum(
            firm_id in selected_pharmacy_firms and row.get("discovery_search_file") == page_file
            for firm_id, row in pharmacy_firms.items()
        )
        search_log.append(
            {
                "query_id": f"2gis_district_pharmacy_page_{index + 1}",
                "query": page["source_url"],
                "scope": page["search_file"],
                "status": page["status"],
                "result_links": page.get("result_links"),
                "new_firm_links_vs_prior_pages": page.get("new_firm_links_vs_prior_pages"),
                "new_verified_physical_places": accepted,
                "result_note": (
                    "Firm IDs are discovery only; selected locations have reviewed "
                    "firm-specific pins."
                ),
                "checked_at": page["checked_at"],
            }
        )
    petgoods_variant_yields = read_jsonl(DATA / "2gis_district_petgoods_search_yields.jsonl")
    for index, page in enumerate(petgoods_variant_yields):
        search_log.append(
            {
                "query_id": f"2gis_district_petgoods_variant_{index + 1}",
                "query": page["source_url"],
                "scope": page["district"],
                "status": page["status"],
                "result_links": page.get("result_links"),
                "new_firm_links_vs_prior_2gis": page.get("new_firm_links_vs_prior_2gis"),
                "new_verified_physical_places": page.get("new_target_physical_places"),
                "result_note": page.get("note"),
                "checked_at": page["checked_at"],
            }
        )
    for probe in probes:
        search_log.append(
            {
                "query_id": "probe_"
                + probe["candidate_id"]
                + "_"
                + ("address" if probe.get("address_hint") else "name"),
                "query": probe["search_url"],
                "scope": probe["candidate_id"],
                "status": "executed",
                "new_candidate_yield": 0,
                "result_note": (
                    "Verification probe only; organization result is not an accepted entity match."
                ),
                "result_url": probe.get("organization_url"),
                "checked_at": probe["checked_at"],
            }
        )
    two_gis_lead_probes = read_jsonl(DATA / "2gis_original_lead_probes.jsonl") + read_jsonl(
        DATA / "2gis_original_lead_probes_short.jsonl"
    )
    for index, probe in enumerate(two_gis_lead_probes):
        search_log.append(
            {
                "query_id": f"2gis_original_lead_probe_{index + 1}",
                "query": probe["source_url"],
                "scope": probe["candidate_id"],
                "status": probe["status"],
                "result_firm_ids": probe["result_firm_ids"],
                "result_count": len(probe["result_firm_ids"]),
                "result_note": (
                    "Search results are suggestions only; exact branch links require "
                    "reviewed firm pages and matching points."
                ),
                "checked_at": probe["checked_at"],
            }
        )

    # Sources are deliberately held separately from user-facing place records.
    write_jsonl(OUT / "places_candidates.jsonl", candidates)
    write_jsonl(OUT / "places_verified.jsonl", verified)
    write_jsonl(OUT / "places_needs_review.jsonl", needs_review)
    write_jsonl(OUT / "places_sources.jsonl", source_rows)
    write_jsonl(OUT / "places_search_log.jsonl", search_log)
    write_jsonl(OUT / "places_audit.jsonl", audit)
    write_jsonl(OUT / "places_other_categories.jsonl", other_categories)
    write_jsonl(OUT / "places_verification_queue.jsonl", triage)

    category_counts = dict(Counter(p["category"] for p in verified))

    def pct(field: str) -> float:
        return (
            round(100 * sum(bool(p[field]) for p in verified) / len(verified), 1)
            if verified
            else 0.0
        )

    original_lead_ids = {row["candidate_id"] for row in probes}
    report = {
        "generated_at": stamp,
        "raw_observations": len(all_observations),
        "total_candidates": len(candidates),
        "verified_physical_places": len(verified),
        "continuation_baseline_candidates": 180,
        "candidates_added_since_baseline": len(candidates) - 180,
        "original_leads_probed": len(original_lead_ids),
        "original_leads_promoted_to_verified": sum(
            row["candidate_id"] in original_lead_ids and row["status"] == "verified"
            for row in candidates
        ),
        "original_leads_still_discovery_only": sum(
            row["candidate_id"] in original_lead_ids and row["status"] == "discovery_only"
            for row in candidates
        ),
        "held_ambiguous_candidate_entities": sum(
            row["status"] == "needs_review" for row in candidates
        ),
        "places_needing_review": len(needs_review),
        "duplicate_observations_rejected": sum(x["type"] == "duplicate_observation" for x in audit),
        "categories_verified": {key: category_counts.get(key, 0) for key in sorted(TARGET)},
        "other_category_or_out_of_scope_observations": len(other_categories),
        "other_categories_discovered": sorted(
            {row["category"] for row in read_jsonl(DATA / "other_category_findings.jsonl")}
        ),
        "field_completeness_pct_verified": {
            "address": pct("address"),
            "coordinates": (
                round(
                    100
                    * sum(
                        p["latitude"] is not None and p["longitude"] is not None for p in verified
                    )
                    / len(verified),
                    1,
                )
                if verified
                else 0.0
            ),
            **{key: pct(key) for key in CONTACT_FIELDS},
        },
        "verified_districts": sorted({p["district"] for p in verified if p["district"]}),
        "district_statistics": {
            district: {
                "candidates": sum(
                    (
                        c.get("district_hint") == district
                        or (
                            c.get("verified_place_id")
                            in {p["id"] for p in verified if p["district"] == district}
                        )
                    )
                    for c in candidates
                ),
                "verified": sum(p["district"] == district for p in verified),
                **{
                    category: sum(
                        p["district"] == district and p["category"] == category for p in verified
                    )
                    for category in sorted(TARGET)
                },
            }
            for district in DISTRICT_HINTS
        },
        "candidates_without_district_hint": sum(not c.get("district_hint") for c in candidates),
        "unresolved_field_conflicts": sum(x["type"] == "field_conflict" for x in audit),
        "curated_places_held_for_review": sum(x["type"] == "curated_record_held" for x in audit),
        "discovery_only_candidates": sum(c["status"] == "discovery_only" for c in candidates),
        "excluded_candidates": sum(c["status"] == "excluded_from_map" for c in candidates),
        "candidate_research_status_counts": dict(Counter(row["research_status"] for row in triage)),
        "public_search_probe_attempts": len(probes),
        "public_search_probed_unique_leads": len({row["candidate_id"] for row in probes}),
        "2gis_original_lead_probe_attempts": len(two_gis_lead_probes),
        "2gis_original_leads_probed_unique": len(
            {row["candidate_id"] for row in two_gis_lead_probes}
        ),
        "2gis_original_leads_with_search_results": len(
            {row["candidate_id"] for row in two_gis_lead_probes if row["result_firm_ids"]}
        ),
        "2gis_original_lead_probe_statuses": dict(
            Counter(row["status"] for row in two_gis_lead_probes)
        ),
        "pagination_yields": [
            {
                "source_url": row["source_url"],
                "page": row["page"],
                "first_seen_name_address_pairs": row["first_seen_name_address_pairs"],
                "target_observations": row["target_observations"],
            }
            for row in pagination_yields
        ],
        "2gis_firm_pages_with_location_suggestions": len(two_gis_suggestions),
        "2gis_individually_selected_physical_places": sum(
            len(read_jsonl(DATA / f"2gis_{brand}_verified_selections.jsonl"))
            for brand in (
                "zooplaneta",
                "animalplanet",
                "arcazoo",
                "felix",
                "neko",
                "district_clinic",
                "district_petstore",
                "district_pharmacy",
            )
        ),
        "2gis_chain_branch_pages_executed": len(two_gis_yields),
        "2gis_district_clinic_pages_executed": len(district_search_yields),
        "2gis_district_petstore_pages_executed": len(petstore_search_yields),
        "2gis_district_petstore_later_pages_executed": sum(
            row["status"] == "executed" for row in petstore_pagination_checks
        ),
        "2gis_district_pharmacy_pages_executed": sum(
            row["status"] == "executed" for row in pharmacy_search_checks
        ),
        "2gis_district_pharmacy_later_pages_executed": sum(
            row["status"] == "executed" for row in pharmacy_pagination_checks
        ),
        "2gis_petgoods_variant_district_pages_executed": sum(
            row["status"] == "executed" for row in petgoods_variant_yields
        ),
        "2gis_petgoods_variant_new_target_places": sum(
            row.get("new_target_physical_places") or 0 for row in petgoods_variant_yields
        ),
        "verified_places_with_field_review": sum(
            r.get("review_status") == "field_conflict" for r in needs_review
        ),
        "source_kinds_used_for_discovery": sorted({o["source_kind"] for o in observations}),
        "source_domains_used_for_verified": sorted(
            {urlparse(u).hostname for p in verified for u in p["source_urls"]}
        ),
        "executed_query_examples_logged": len(read_jsonl(DATA / "executed_queries.jsonl")),
        "continuation_queries_logged": len(read_jsonl(DATA / "continuation_queries.jsonl")),
        "district_sweep_queries_logged": len(read_jsonl(DATA / "district_sweep_queries.jsonl")),
        "continuation_page_yields": {
            label: source_stats.get(url, {}).get("first_seen_unique", 0)
            for label, url in {
                "yandex_veterinary_pharmacy_page_2": "https://yandex.com/maps/10335/tashkent/search/%D0%92%D0%B5%D1%82%D0%B0%D0%BF%D1%82%D0%B5%D0%BA%D0%B0/?page=2",
                "yandex_veterinary_clinic_page_2": "https://yandex.com/maps/10335/tashkent/category/veterinary_clinic/184107216/?page=2",
                "yandex_veterinary_clinic_page_3": "https://yandex.com/maps/10335/tashkent/category/veterinary_clinic/184107216/?page=3",
                "yandex_pet_goods_page_6": "https://yandex.com/maps/10335/tashkent/search/%D0%A2%D0%BE%D0%B2%D0%B0%D1%80%D1%8B%20%D0%B4%D0%BB%D1%8F%20%D0%B6%D0%B8%D0%B2%D0%BE%D1%82%D0%BD%D1%8B%D1%85/?page=6",
            }.items()
        },
        "pet_shop_page_first_seen_candidates": {
            str(page): source_stats.get(
                f"https://yandex.com/maps/10335/tashkent/category/pet_shop/184107222/?page={page}",
                {},
            ).get("first_seen_unique", 0)
            for page in range(2, 6)
        },
        "saturation": "not_demonstrated",
        "saturation_reason": (
            "Accessible pet-goods pages 7-10 and veterinary-pharmacy pages 3-4 continued to yield "
            "first-seen name/address pairs. The next page in each sequence returned no "
            "target listings, but later district 2GIS pet-store pages and veterinary-pharmacy "
            "searches still found new physical places. A separate 12-district 2GIS pet-goods "
            "wording sweep yielded zero target physical places after firm-page review. One "
            "Yangihayot pharmacy search returned HTTP 404; citywide saturation remains unproven."
        ),
        "production_changed": False,
    }
    (OUT / "quality_report.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    make_workbook(candidates, verified, needs_review, source_rows, audit, report)
    return report


def make_workbook(
    candidates: list[dict],
    verified: list[dict],
    review: list[dict],
    sources: list[dict],
    audit: list[dict],
    report: dict,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    sheets = [
        (
            "Verified",
            verified,
            [
                "id",
                "name",
                "branch_name",
                "category",
                "district",
                "address",
                "latitude",
                "longitude",
                "phones",
                "website",
                "instagram",
                "telegram",
                "opening_hours",
                "review_status",
                "entity_confidence",
                "field_confidence",
                "source_count",
                "missing_phone",
                "missing_hours",
                "field_conflicts",
                "source_urls",
            ],
        ),
        (
            "Needs review",
            review,
            [
                "id",
                "name",
                "name_hint",
                "category",
                "category_hint",
                "address",
                "address_hint",
                "district",
                "district_hint",
                "review_status",
                "entity_confidence",
                "source_count",
                "missing_phone",
                "missing_hours",
                "field_conflicts",
                "review_reasons",
                "source_urls",
            ],
        ),
        (
            "Candidates",
            candidates,
            [
                "candidate_id",
                "name_hint",
                "category_hint",
                "address_hint",
                "district_hint",
                "status",
                "source_count",
                "verified_place_id",
                "observation_ids",
                "source_urls",
            ],
        ),
        (
            "Verification queue",
            read_jsonl(OUT / "places_verification_queue.jsonl"),
            [
                "candidate_id",
                "name_hint",
                "category_hint",
                "district_hint",
                "status",
                "research_status",
                "source_count",
                "source_kinds",
                "next_action",
                "note",
            ],
        ),
        (
            "Sources",
            sources,
            [
                "record_id",
                "field",
                "value",
                "source_url",
                "source_kind",
                "confidence",
                "status",
                "branch_scope",
                "checked_at",
            ],
        ),
        (
            "Audit",
            audit,
            [
                "type",
                "place_id",
                "place_ids",
                "candidate_id",
                "candidate_ids",
                "observation_id",
                "reason",
                "reasons",
                "evidence_urls",
            ],
        ),
    ]
    for title, rows, columns in sheets:
        ws = wb.create_sheet(title)
        ws.append(columns)
        for row in rows:
            ws.append(
                [
                    (
                        json.dumps(row.get(key), ensure_ascii=False)
                        if isinstance(row.get(key), (dict, list))
                        else row.get(key)
                    )
                    for key in columns
                ]
            )
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="23405A")
        for col in ws.columns:
            width = min(75, max(14, max(len(str(c.value or "")) for c in list(col)[:100]) + 2))
            ws.column_dimensions[get_column_letter(col[0].column)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if title == "Sources" and cell.column == 4 and isinstance(cell.value, str):
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
    summary = wb.create_sheet("Summary", 0)
    summary.append(["Metric", "Value"])
    for key, value in report.items():
        summary.append(
            [
                key,
                json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value,
            ]
        )
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 100
    wb.save(OUT / "places_review.xlsx")


if __name__ == "__main__":
    print(json.dumps(build(), indent=2, ensure_ascii=False))
