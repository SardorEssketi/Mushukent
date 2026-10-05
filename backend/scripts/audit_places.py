"""Build conservative, reviewable place datasets without accessing a database.

Curated evidence is intentionally separate from the legacy workbook. A value is
publishable only when an evidence item names its source and confidence.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path
from urllib.parse import urlparse
from uuid import NAMESPACE_URL, uuid5

FIELDS = (
    "name",
    "category",
    "subcategory",
    "latitude",
    "longitude",
    "address",
    "district",
    "phone_numbers",
    "website",
    "instagram",
    "telegram",
    "other_contacts",
    "opening_hours",
    "description",
)
LEGACY_FIELDS = {
    "name": "name",
    "category": "type",
    "latitude": "latitude",
    "longitude": "longitude",
    "address": "address",
    "website": "website",
    "instagram": "instagram",
    "telegram": "telegram",
    "opening_hours": "opening_hours",
    "description": "description",
}
TASHKENT_BOUNDS = (41.18, 41.43, 69.05, 69.42)
CANONICAL_CATEGORIES = {"pet_store", "veterinary_clinic", "veterinary_pharmacy", "animal_shelter"}
PHONE_RE = re.compile(r"^\+998\d{9}$")


def stable_id(kind: str, key: str) -> str:
    return str(uuid5(NAMESPACE_URL, f"mushukistan-place-audit:{kind}:{key}"))


def normalized_phone(value: str) -> str | None:
    digits = re.sub(r"\D", "", value)
    candidate = f"+{digits}"
    return candidate if PHONE_RE.fullmatch(candidate) else None


def valid_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme in {"https", "http"} and bool(parsed.hostname)


def distance_m(a: dict, b: dict) -> float:
    lat1, lon1 = math.radians(a["latitude"]), math.radians(a["longitude"])
    lat2, lon2 = math.radians(b["latitude"]), math.radians(b["longitude"])
    value = (
        math.sin((lat2 - lat1) / 2) ** 2
        + math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2
    )
    return 12742000 * math.asin(min(1, math.sqrt(value)))


def read_workbook(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    expected = {"name", "type", "latitude", "longitude", "address", "phone", "phone 2", "website"}
    if not expected.issubset(set(headers)) or len(headers) != len(set(headers)):
        raise ValueError("Workbook columns changed or contain duplicates")
    result = []
    for row_number, values in enumerate(rows, 2):
        if all(value is None for value in values):
            continue
        row = {header: value for header, value in zip(headers, values, strict=True)}
        row["row_number"] = row_number
        row["place_id"] = stable_id("legacy-row", str(row_number))
        result.append(row)
    return result


def write_jsonl(path: Path, rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def add_flag(flags: list[dict], row: dict, code: str, detail: str) -> None:
    flags.append(
        {
            "place_id": row["place_id"],
            "row_number": row.get("row_number"),
            "code": code,
            "detail": detail,
        }
    )


def check_legacy(rows: list[dict]) -> list[dict]:
    flags: list[dict] = []
    phones: dict[str, list[dict]] = defaultdict(list)
    domains: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        for field in ("latitude", "longitude"):
            if not isinstance(row[field], (int, float)) or not math.isfinite(row[field]):
                add_flag(flags, row, "invalid_coordinate", field)
        if isinstance(row["latitude"], (int, float)) and isinstance(row["longitude"], (int, float)):
            south, north, west, east = TASHKENT_BOUNDS
            if not (south <= row["latitude"] <= north and west <= row["longitude"] <= east):
                add_flag(
                    flags, row, "outside_tashkent_bbox", "Coordinate is outside the import bbox"
                )
        for field in ("phone", "phone 2"):
            if row[field]:
                phone = normalized_phone(str(row[field]))
                if phone is None:
                    add_flag(flags, row, "invalid_uz_phone", f"{field}: {row[field]}")
                else:
                    phones[phone].append(row)
        for field in ("website", "instagram", "telegram"):
            value = row.get(field)
            if not value:
                continue
            url = str(value) if "://" in str(value) else f"https://{value}"
            if not valid_url(url):
                add_flag(flags, row, "invalid_url", f"{field}: {value}")
            elif field == "website":
                domains[urlparse(url).hostname.removeprefix("www.")].append(row)
    for phone, group in phones.items():
        if len({r["name"] for r in group}) > 1:
            for row in group:
                add_flag(flags, row, "shared_phone_across_names", phone)
    for domain, group in domains.items():
        if len({r["name"] for r in group}) > 1:
            for row in group:
                add_flag(flags, row, "shared_website_across_names", domain)
    for i, left in enumerate(rows):
        for right in rows[i + 1 :]:
            if not all(
                isinstance(r[k], (float, int))
                for r in (left, right)
                for k in ("latitude", "longitude")
            ):
                continue
            meters = distance_m(left, right)
            same_name = re.sub(r"\W+", "", left["name"].casefold()) == re.sub(
                r"\W+", "", right["name"].casefold()
            )
            left_phones = {
                normalized_phone(str(left[key])) for key in ("phone", "phone 2") if left[key]
            }
            right_phones = {
                normalized_phone(str(right[key])) for key in ("phone", "phone 2") if right[key]
            }
            common_phone = bool((left_phones & right_phones) - {None})
            same_address = bool(
                left.get("address")
                and right.get("address")
                and re.sub(r"\W+", "", str(left["address"]).casefold())
                == re.sub(r"\W+", "", str(right["address"]).casefold())
            )
            if meters < 10 or (meters < 80 and same_name and (common_phone or same_address)):
                code = "possible_duplicate" if same_name else "shared_site_distinct_businesses"
                add_flag(flags, left, code, f"row {right['row_number']} is {meters:.1f} m away")
                add_flag(flags, right, code, f"row {left['row_number']} is {meters:.1f} m away")
    return flags


def validate_evidence(item: dict) -> None:
    if item["confidence"] not in {"high", "medium", "low", "conflict"}:
        raise ValueError("Invalid evidence confidence")
    if not item.get("sources") or not all(valid_url(url) for url in item["sources"]):
        raise ValueError("Each evidence item needs valid source URLs")


def build_record(
    raw: dict, evidence: dict | None, flags: list[dict], timestamp: str
) -> tuple[dict, list[dict]]:
    is_new = "row_number" not in raw
    record = {
        "place_id": raw["place_id"],
        "origin": "new" if is_new else "legacy_workbook",
        "legacy_row": raw.get("row_number"),
        "alternative_names": [],
        "source_references": [],
        "last_verified_at": None,
        "evidence_checked_at": None,
        "review_status": "needs_review",
        "review_reasons": [],
        "field_evidence": {},
    }
    record.update({field: None for field in FIELDS})
    changes = []
    record["candidate"] = {
        "name": raw.get("name") if not is_new else None,
        "latitude": raw.get("latitude") if not is_new else None,
        "longitude": raw.get("longitude") if not is_new else None,
        "address": raw.get("address") if not is_new else None,
    }
    if evidence:
        for field, item in evidence.get("fields", {}).items():
            if field not in FIELDS:
                raise ValueError(f"Unsupported field: {field}")
            validate_evidence(item)
            if field == "category" and item["value"] not in CANONICAL_CATEGORIES:
                raise ValueError(f"Unsupported canonical category: {item['value']}")
            record["field_evidence"][field] = item
            record["source_references"].extend(item["sources"])
            if item["confidence"] == "high":
                record[field] = item["value"]
            else:
                record["review_reasons"].append(f"{field}: {item['confidence']}")
                if field in record["candidate"]:
                    record["candidate"][field] = item["value"]
        record["source_references"] = sorted(set(record["source_references"]))
        record["alternative_names"] = evidence.get("alternative_names", [])
        if record["source_references"]:
            record["evidence_checked_at"] = timestamp
    for field in FIELDS:
        old = raw.get(LEGACY_FIELDS[field]) if field in LEGACY_FIELDS and not is_new else None
        if field == "phone_numbers" and not is_new:
            old = [str(raw[k]) for k in ("phone", "phone 2") if raw[k]] or None
        new = record[field]
        if old == new or old is None and new is None:
            continue
        evidence_item = record["field_evidence"].get(field)
        changes.append(
            {
                "place_id": record["place_id"],
                "legacy_row": raw.get("row_number"),
                "field": field,
                "old_value": old,
                "new_value": new,
                "sources": evidence_item["sources"] if evidence_item else [],
                "reason": (
                    "credible sources conflict; value withheld"
                    if evidence_item and evidence_item["confidence"] == "conflict"
                    else (
                        "supported by cited evidence"
                        if evidence_item and new is not None
                        else "legacy value withheld pending independent verification"
                    )
                ),
                "confidence": (evidence_item["confidence"] if evidence_item else "unverified"),
            }
        )
    essential = ("name", "category", "latitude", "longitude", "address")
    has_conflict = any(
        item["confidence"] == "conflict" for item in record["field_evidence"].values()
    )
    if all(record[field] is not None for field in essential) and not has_conflict:
        record["review_status"] = "high_confidence"
    elif not all(record[field] is not None for field in essential):
        record["review_reasons"].append(
            "identity, address, or coordinates lack high-confidence evidence"
        )
    if record["latitude"] is not None and record["longitude"] is not None:
        south, north, west, east = TASHKENT_BOUNDS
        if not (south <= record["latitude"] <= north and west <= record["longitude"] <= east):
            record["review_status"] = "needs_review"
            record["review_reasons"].append("verified coordinates outside Tashkent bbox")
        if (
            not is_new
            and isinstance(raw.get("latitude"), (int, float))
            and isinstance(raw.get("longitude"), (int, float))
        ):
            meters = distance_m(raw, record)
            if meters > 100:
                record["review_status"] = "needs_review"
                record["review_reasons"].append(
                    f"coordinate moved {meters:.0f} m from legacy candidate"
                )
    if record["phone_numbers"] is not None and any(
        normalized_phone(phone) is None for phone in record["phone_numbers"]
    ):
        raise ValueError(f"Evidence has invalid phone: {record['place_id']}")
    if record["review_status"] == "high_confidence":
        record["last_verified_at"] = timestamp
    for field in ("website", "instagram", "telegram"):
        if record[field] is not None and not valid_url(record[field]):
            raise ValueError(f"Evidence has invalid URL: {field}")
    if record["instagram"] and urlparse(record["instagram"]).hostname not in {
        "instagram.com",
        "www.instagram.com",
    }:
        raise ValueError("Instagram evidence has an unexpected host")
    if record["telegram"] and urlparse(record["telegram"]).hostname not in {"t.me", "telegram.me"}:
        raise ValueError("Telegram evidence has an unexpected host")
    own_flags = [flag for flag in flags if flag["place_id"] == record["place_id"]]
    record["validation_flags"] = own_flags
    if not is_new:
        record["legacy_candidate"] = {k: v for k, v in raw.items() if k not in {"place_id"}}
    return record, changes


def write_workbook(
    path: Path, reviewed: list[dict], pending: list[dict], new: list[dict], flags: list[dict]
) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in (("Reviewed", reviewed), ("Needs review", pending), ("New candidates", new)):
        sheet = workbook.create_sheet(title)
        sheet.append(
            [
                "place_id",
                "legacy_row",
                "name",
                "candidate_name",
                "category",
                "latitude",
                "longitude",
                "candidate_address",
                "address",
                "phone_numbers",
                "website",
                "opening_hours",
                "review_status",
                "review_reasons",
                "source_references",
            ]
        )
        for row in rows:
            values = [
                row.get("place_id"),
                row.get("legacy_row"),
                row.get("name"),
                row["candidate"].get("name"),
                row.get("category"),
                row.get("latitude") or row["candidate"].get("latitude"),
                row.get("longitude") or row["candidate"].get("longitude"),
                row["candidate"].get("address"),
                row.get("address"),
                row.get("phone_numbers"),
                row.get("website"),
                row.get("opening_hours"),
                row.get("review_status"),
                row.get("review_reasons"),
                row.get("source_references"),
            ]
            sheet.append(
                [
                    (
                        json.dumps(value, ensure_ascii=False)
                        if isinstance(value, (list, dict))
                        else value
                    )
                    for value in values
                ]
            )
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = sheet.dimensions
    sheet = workbook.create_sheet("Validation flags")
    sheet.append(["place_id", "row_number", "code", "detail"])
    for flag in flags:
        sheet.append([flag[key] for key in ("place_id", "row_number", "code", "detail")])
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=Path("Mushukistan_map.xlsx"))
    parser.add_argument("--evidence", type=Path, default=Path("data/places_review/evidence.json"))
    parser.add_argument("--output", type=Path, default=Path("data/places_review/output"))
    args = parser.parse_args()
    evidence = json.loads(args.evidence.read_text(encoding="utf-8"))
    workbook_hash = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    if workbook_hash != evidence["source_sha256"]:
        raise ValueError(
            "Source workbook changed; reconcile row IDs and evidence before rebuilding"
        )
    raw = read_workbook(args.workbook)
    existing_by_row = {entry["row_number"]: entry for entry in evidence["existing"]}
    if len(existing_by_row) != len(evidence["existing"]):
        raise ValueError("Evidence contains duplicate workbook row numbers")
    if set(existing_by_row) - {row["row_number"] for row in raw}:
        raise ValueError("Evidence references absent workbook rows")
    flags = check_legacy(raw)
    timestamp = datetime.now(UTC).isoformat()
    records, audit = [], []
    for row in raw:
        record, changes = build_record(
            row, existing_by_row.get(row["row_number"]), flags, timestamp
        )
        records.append(record)
        audit.extend(changes)
    new = []
    for entry in evidence["new"]:
        new_id = stable_id("new", entry["candidate_key"])
        record, changes = build_record({"place_id": new_id}, entry, flags, timestamp)
        new.append(record)
        audit.extend(changes)
    reviewed = [record for record in records + new if record["review_status"] == "high_confidence"]
    pending = [record for record in records + new if record["review_status"] != "high_confidence"]
    ids = [record["place_id"] for record in records + new]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate stable IDs")
    output = args.output
    output.mkdir(parents=True, exist_ok=True)
    write_jsonl(output / "places_raw_snapshot.jsonl", raw)
    write_jsonl(output / "places_enriched_reviewed.jsonl", reviewed)
    write_jsonl(output / "places_needs_review.jsonl", pending)
    write_jsonl(output / "places_new_candidates.jsonl", new)
    write_jsonl(output / "places_audit_log.jsonl", audit)
    write_jsonl(output / "places_validation_flags.jsonl", flags)
    write_workbook(output / "places_review.xlsx", reviewed, pending, new, flags)
    coord_rows = set()
    address_rows = set()
    phone_rows = set()
    phone_enriched_rows = set()
    website_rows = set()
    for row in raw:
        current = next(record for record in records if record["legacy_row"] == row["row_number"])
        if current["latitude"] is not None and distance_m(row, current) > 10:
            coord_rows.add(row["row_number"])
        if (
            "address" in existing_by_row.get(row["row_number"], {}).get("material_corrections", [])
            and current["address"] is not None
        ):
            address_rows.add(row["row_number"])
        if current["phone_numbers"] is not None:
            old_phones = [
                normalized_phone(str(row[key])) for key in ("phone", "phone 2") if row[key]
            ]
            if "phone_numbers" in existing_by_row.get(row["row_number"], {}).get(
                "material_corrections", []
            ):
                phone_rows.add(row["row_number"])
            elif any(phone not in old_phones for phone in current["phone_numbers"]):
                phone_enriched_rows.add(row["row_number"])
        if "website" in existing_by_row.get(row["row_number"], {}).get("material_corrections", []):
            website_rows.add(row["row_number"])
    completeness = {
        field: {
            "count": sum(record[field] is not None for record in reviewed),
            "percent": (
                round(
                    100 * sum(record[field] is not None for record in reviewed) / len(reviewed), 1
                )
                if reviewed
                else 0
            ),
        }
        for field in FIELDS
    }
    legacy_completeness = {
        label: {
            "count": sum(bool(row.get(column)) for row in raw),
            "percent": round(100 * sum(bool(row.get(column)) for row in raw) / len(raw), 1),
        }
        for label, column in (
            ("coordinates", "latitude"),
            ("address", "address"),
            ("phone", "phone"),
            ("website", "website"),
            ("instagram", "instagram"),
            ("telegram", "telegram"),
            ("opening_hours", "opening_hours"),
        )
    }
    summary = {
        "existing_places_examined": len(records),
        "existing_places_with_any_cited_evidence": sum(
            bool(record["source_references"]) for record in records
        ),
        "existing_places_high_confidence": sum(
            record["review_status"] == "high_confidence" for record in records
        ),
        "new_places_discovered": len(new),
        "new_places_high_confidence": sum(
            record["review_status"] == "high_confidence" for record in new
        ),
        "final_high_confidence_places": len(reviewed),
        "existing_places_materially_corrected": len(
            coord_rows | address_rows | phone_rows | website_rows
        ),
        "corrected_phone_records": len(phone_rows),
        "enriched_phone_records": len(phone_enriched_rows),
        "corrected_website_records": len(website_rows),
        "corrected_address_records": len(address_rows),
        "corrected_coordinate_records_over_10m": len(coord_rows),
        "unresolved_review_records": len(pending),
        "unresolved_conflict_records": sum(
            any(item["confidence"] == "conflict" for item in record["field_evidence"].values())
            for record in records + new
        ),
        "reviewed_official_websites": sum(record["website"] is not None for record in reviewed),
        "reviewed_instagram": sum(record["instagram"] is not None for record in reviewed),
        "reviewed_telegram": sum(record["telegram"] is not None for record in reviewed),
        "reviewed_opening_hours": sum(record["opening_hours"] is not None for record in reviewed),
        "possible_duplicate_flags": sum(flag["code"] == "possible_duplicate" for flag in flags)
        // 2,
        "validation_flag_counts": dict(Counter(flag["code"] for flag in flags)),
        "field_completeness_reviewed": completeness,
        "legacy_raw_completeness_unverified": legacy_completeness,
        "generated_at": timestamp,
    }
    (output / "quality_report.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(json.dumps(summary, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
