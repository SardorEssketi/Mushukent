"""Build a human review workbook from the immutable production comparison."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "places_v2" / "output"
WORKBOOK = OUT / "production_comparison_review.xlsx"


def read_jsonl(name: str) -> list[dict]:
    return [json.loads(line) for line in (OUT / name).read_text(encoding="utf-8").splitlines()]


def add_sheet(book: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    page = book.create_sheet(title)
    page.append(headers)
    for row in rows:
        page.append(row)
    page.freeze_panes = "C2"
    page.auto_filter.ref = page.dimensions
    for cell in page[1]:
        cell.fill = PatternFill("solid", fgColor="244268")
        cell.font = Font(color="FFFFFF", bold=True)
    page.column_dimensions["A"].width = 38
    page.column_dimensions["B"].width = 30
    page.column_dimensions["C"].width = 32


def main() -> None:
    if WORKBOOK.exists():
        raise FileExistsError(f"Refusing to overwrite review workbook: {WORKBOOK}")
    diff = read_jsonl("production_comparison.jsonl")
    review = read_jsonl("production_ambiguous_review.jsonl")
    report = json.loads((OUT / "production_comparison_report.json").read_text(encoding="utf-8"))
    book = Workbook()
    book.remove(book.active)
    add_sheet(
        book,
        "Summary",
        ["metric", "value"],
        [
            [key, json.dumps(value, ensure_ascii=False) if isinstance(value, dict) else value]
            for key, value in report.items()
        ],
    )
    matched = [row for row in diff if row["action"] == "UPDATE_EXISTING"]
    add_sheet(
        book,
        "Matched updates",
        [
            "production_uuid",
            "production_name",
            "v2_name",
            "v2_id",
            "distance_m",
            "name_similarity",
            "address_similarity",
            "production_category",
            "v2_category",
            "production_address",
            "v2_address",
            "production_latitude",
            "production_longitude",
            "v2_latitude",
            "v2_longitude",
            "old_source_id",
            "new_source_id",
            "legacy_fields_to_clear",
            "proposed_changes_json",
            "human_approve",
        ],
        [
            [
                row["production_uuid"],
                row["production"]["name"],
                row["v2"]["name"],
                row["v2_id"],
                row["match_evidence"]["distance_m"],
                row["match_evidence"]["name_similarity"],
                row["match_evidence"]["address_similarity"],
                row["production"]["category"],
                row["v2"]["category"],
                row["production"]["address"],
                row["v2"]["address"],
                row["production"]["latitude"],
                row["production"]["longitude"],
                row["v2"]["latitude"],
                row["v2"]["longitude"],
                row["production"]["source_id"],
                row["attach_source_id"],
                ", ".join(row["legacy_nonnull_to_v2_null"]),
                json.dumps(row["proposed_field_changes"], ensure_ascii=False),
                None,
            ]
            for row in matched
        ],
    )
    add_sheet(
        book,
        "Production only",
        [
            "production_uuid",
            "name",
            "category",
            "address",
            "latitude",
            "longitude",
            "phone",
            "website",
            "opening_hours",
            "source_id",
            "created_at",
            "decision_notes",
        ],
        [
            [
                row["production_uuid"],
                row["production"]["name"],
                row["production"]["category"],
                row["production"]["address"],
                row["production"]["latitude"],
                row["production"]["longitude"],
                row["production"]["phone"],
                row["production"]["website"],
                row["production"]["opening_hours"],
                row["production"]["source_id"],
                row["production"]["created_at"],
                None,
            ]
            for row in diff
            if row["action"] == "KEEP_UNCHANGED"
        ],
    )
    add_sheet(
        book,
        "Ambiguous pairs",
        [
            "production_uuid",
            "production_name",
            "production_address",
            "production_latitude",
            "production_longitude",
            "v2_id",
            "v2_name",
            "v2_address",
            "v2_latitude",
            "v2_longitude",
            "distance_m",
            "name_similarity",
            "address_similarity",
            "category_agreement",
            "phone_agreement",
            "website_agreement",
            "reason",
            "decision_notes",
        ],
        [
            [
                row["production_uuid"],
                row["production_name"],
                row["production_address"],
                row["production_latitude"],
                row["production_longitude"],
                row["v2_id"],
                row["v2_name"],
                row["v2_address"],
                row["v2_latitude"],
                row["v2_longitude"],
                row["distance_m"],
                row["matching_evidence"]["name_similarity"],
                row["matching_evidence"]["address_similarity"],
                row["matching_evidence"]["category_agreement"],
                row["matching_evidence"]["phone_agreement"],
                row["matching_evidence"]["website_agreement"],
                row["reason"],
                None,
            ]
            for row in review
        ],
    )
    for action, title in (("INSERT_NEW", "V2 proposed inserts"), ("REVIEW", "V2 held for review")):
        add_sheet(
            book,
            title,
            [
                "v2_id",
                "v2_name",
                "v2_category",
                "v2_address",
                "latitude",
                "longitude",
                "stable_source_id",
                "candidate_production_uuids",
                "source_urls",
                "decision_notes",
            ],
            [
                [
                    row["v2_id"],
                    row["v2"]["name"],
                    row["v2"]["category"],
                    row["v2"]["address"],
                    row["v2"]["latitude"],
                    row["v2"]["longitude"],
                    row["v2"]["source_id"],
                    ", ".join(row.get("candidate_production_uuids", [])),
                    "\n".join(row["v2"]["source_urls"]),
                    None,
                ]
                for row in diff
                if row["entity"] == "v2" and row["action"] == action
            ],
        )
    book.save(WORKBOOK)
    print(
        json.dumps(
            {
                "workbook": str(WORKBOOK),
                "matched": len(matched),
                "production_only": sum(row["action"] == "KEEP_UNCHANGED" for row in diff),
                "ambiguous_pairs": len(review),
            }
        )
    )


if __name__ == "__main__":
    main()
